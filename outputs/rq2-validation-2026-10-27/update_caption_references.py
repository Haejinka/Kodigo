from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
import re,json
R=Path.cwd();B=R/'outputs/rq2-validation-2026-10-27';P=R/'deliverables/thesis/Chapter 1-4 Manuscript v3 RQ2 Technical Review.docx'
d=Document(P)
# Lists in this manuscript are native tables. Add the two new evidence tables.
t=d.tables[1]
for number,title in [('4.6','Executed Technical Verification Results'),('4.7','Confirmed Authorization Failures in the Controlled Assessment')]:
 if not any(row.cells[0].text.strip()==number for row in t.rows):
  row=deepcopy(t.rows[-1]._tr);t._tbl.append(row)
  cells=t.rows[-1].cells;cells[0].text=number;cells[1].text=title;cells[2].text=''
  for cell in cells:
   for p in cell.paragraphs:
    for r in p.runs:r.font.name='Cambria';r.font.size=Pt(11)
ids=[int(x.get(qn('w:id'))) for x in d.element.findall('.//'+qn('w:bookmarkStart')) if x.get(qn('w:id'),'').isdigit()]
bookmark_id=max(ids+[0])+1;mapping=[]
for table_index,kind in [(0,'Figure'),(1,'Table')]:
 for row in d.tables[table_index].rows[1:]:
  number=row.cells[0].text.strip();needle=f'{kind} {number}.'
  caption=next((p for p in d.paragraphs if p.text.strip().startswith(needle)),None)
  if caption is None:
   caption=next((p for p in d.paragraphs if re.match(r'^'+kind+r'\s+'+re.escape(number)+r'\b',p.text.strip()) and len(p.text)<260),None)
  if caption is None:raise RuntimeError('Caption not found: '+needle)
  name='kodigo_'+kind.lower()+'_'+number.replace('.','_')
  if not d.element.xpath('.//w:bookmarkStart[@w:name="'+name+'"]'):
   start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bookmark_id));start.set(qn('w:name'),name)
   end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark_id));bookmark_id+=1
   caption._p.insert(1 if caption._p.pPr is not None else 0,start);caption._p.append(end)
  pg=row.cells[2].paragraphs[0];pg.clear()
  r=OxmlElement('w:r');fld=OxmlElement('w:fldChar');fld.set(qn('w:fldCharType'),'begin');r.append(fld);pg._p.append(r)
  r=OxmlElement('w:r');instr=OxmlElement('w:instrText');instr.set(qn('xml:space'),'preserve');instr.text=' PAGEREF '+name+' \\h ';r.append(instr);pg._p.append(r)
  r=OxmlElement('w:r');fld=OxmlElement('w:fldChar');fld.set(qn('w:fldCharType'),'separate');r.append(fld);pg._p.append(r)
  r=OxmlElement('w:r');txt=OxmlElement('w:t');txt.text='0';r.append(txt);pg._p.append(r)
  r=OxmlElement('w:r');fld=OxmlElement('w:fldChar');fld.set(qn('w:fldCharType'),'end');r.append(fld);pg._p.append(r)
  mapping.append({'kind':kind,'number':number,'bookmark':name,'caption':caption.text})
d.save(P);(B/'caption-reference-map.json').write_text(json.dumps(mapping,indent=2),encoding='utf8');print('Updated',len(mapping),'dynamic caption page references')

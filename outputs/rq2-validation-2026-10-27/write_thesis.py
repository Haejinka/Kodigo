from pathlib import Path
from copy import deepcopy
from zipfile import ZipFile
from lxml import etree as E
import json,re,hashlib,shutil
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(r'C:\Kodigo v0.1.0');B=ROOT/'outputs/rq2-validation-2026-10-27';O=ROOT/'deliverables/thesis'
O.mkdir(exist_ok=True)
SOURCE=ROOT/'docs/Chapter 1-4 Manuscript v3.docx'
TARGET=O/'Chapter 1-4 Manuscript v3 RQ2 Technical Review.docx'
shutil.copy2(SOURCE,B/'manuscript-before-rq2.docx')
def unpack(file):
 x=json.loads((B/file).read_text(encoding='utf8'))
 if 'content' not in x:return x
 x=json.loads(next(c['text'] for c in x['content'] if c['type']=='text'))
 if isinstance(x.get('result'),str):return json.loads(re.search(r'\n<untrusted-data-[^>]+>\s*([\s\S]*?)\s*</untrusted-data-',x['result']).group(1))
 return x
db=unpack('controlled-security-results-v2.json')[0]['payload'];tests=db['results']
assert len(tests)==22 and sum(t['passed'] for t in tests)==20 and db['fixture_transaction_rolled_back']
local=unpack('local-checks.json');network=unpack('network-and-client-results.json')
assert all(x['exitCode']==0 for x in local['results'])
assert all(x['passed'] for x in network['network']+network['browser'])
metadata=unpack('live-metadata.json')[0]['metadata']
assert len(metadata['tables'])==36 and all(t['rls_enabled'] for t in metadata['tables'])
assert len(metadata['views'])==8 and all('security_invoker=true' in v['reloptions'] for v in metadata['views'])
advisors=unpack('live-initial-0.json')['result']['lints']
d=Document(SOURCE)
original_text=[p.text for p in d.paragraphs]
body_template=next(p for p in d.paragraphs if p.text.startswith('This chapter presents, analyzes'))
head_template=next(p for p in d.paragraphs if p.text=='Summary of Findings')

def font(run,size=11,bold=False):
 run.font.name='Cambria';run.font.size=Pt(size);run.font.bold=bold;run.font.color.rgb=RGBColor(0,0,0)
 run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'),'Cambria')
def para(anchor,txt,heading=False):
 p=anchor.insert_paragraph_before()
 p.style='Heading 1' if heading else 'Normal'
 if not heading and body_template._p.pPr is not None:
  p._p.remove(p._p.pPr)
  pp=deepcopy(body_template._p.pPr)
  for e in list(pp):
   if e.tag.endswith('pPrChange'):pp.remove(e)
  p._p.insert(0,pp)
 if heading:
  p.paragraph_format.keep_with_next=True;p.paragraph_format.space_before=Pt(12);p.paragraph_format.space_after=Pt(6)
  p.paragraph_format.first_line_indent=Inches(0);p.paragraph_format.line_spacing=1.25
 else:
  p.paragraph_format.line_spacing=2;p.paragraph_format.space_after=Pt(0)
  p.paragraph_format.first_line_indent=Inches(.5);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
 font(p.add_run(txt),11,heading)
 return p
def replace(p,txt):
 p.clear();font(p.add_run(txt))
def table(anchor,caption,headers,rows,widths):
 cp=para(anchor,caption);cp.paragraph_format.first_line_indent=0;cp.paragraph_format.line_spacing=1.2;cp.paragraph_format.space_before=Pt(8);cp.paragraph_format.space_after=Pt(6);cp.paragraph_format.keep_with_next=True
 font(cp.runs[0],11,True)
 t=d.add_table(rows=1,cols=len(headers));anchor._p.addprevious(t._tbl)
 t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 for j,w in enumerate(widths):t.columns[j].width=Inches(w)
 for j,h in enumerate(headers):t.rows[0].cells[j].text=h
 for row in rows:
  cells=t.add_row().cells
  for j,v in enumerate(row):cells[j].text=str(v)
 for ri,row in enumerate(t.rows):
  row.height=None
  trPr=row._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');trPr.append(cant)
  if ri==0:
   rep=OxmlElement('w:tblHeader');rep.set(qn('w:val'),'true');trPr.append(rep)
  for j,cell in enumerate(row.cells):
   cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   pr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'5');e.set(qn('w:color'),'D9D9D9');borders.append(e)
   pr.append(borders);marg=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    e=OxmlElement('w:'+edge);e.set(qn('w:w'),'85');e.set(qn('w:type'),'dxa');marg.append(e)
   pr.append(marg)
   if ri==0:
    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7E6E6');pr.append(sh)
   for p in cell.paragraphs:
    p.paragraph_format.first_line_indent=0;p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j==1 else WD_ALIGN_PARAGRAPH.LEFT
    for r in p.runs:font(r,10,ri==0)
 return t

# Clarify how executed verification relates to the retained research design.
survey=next(p for p in d.paragraphs if p.text.startswith('Technical Validation Survey.'))
after=survey._p.getnext()
from docx.text.paragraph import Paragraph
method_anchor=Paragraph(after,survey._parent)
para(method_anchor,'Implementation Verification Supplement',True)
para(method_anchor,'An AI-assisted technical assessment was conducted on 27 October 2026 to examine the implemented framework and information security measures against the four criteria of Research Question 2. Evidence consisted of source-code inspection, local compilation and unit tests, browser route checks, live database configuration, controlled database scenarios, HTTPS requests, and the Supabase Security Advisor. Expected behavior was specified for each scenario before interpreting its result. The assessment used the labels passed, failed, inspected, and not tested; it did not convert technical outcomes into questionnaire ratings.')
para(method_anchor,'The controlled database checks used generated test identities and records within a subtransaction that was deliberately rolled back. PostgreSQL roles and JWT claims simulated authorization contexts; these checks did not issue Auth sessions or perform a real authenticator-code challenge. The expert survey described in Table 3.1 remains a separate evaluation strand. Its respondent counts, mean ratings, and standard deviations are not supplied by this implementation-verification exercise. The evidence therefore supports a technical assessment of RQ2 rather than a claim of completed expert-survey validation.')

# Broaden Chapter 4 scope without changing the RQ1 evidence or inventing RQ3 results.
intro=next(p for p in d.paragraphs if p.text.startswith('This chapter presents, analyzes'))
replace(intro,'This chapter presents the operational findings for Research Question 1 and the implementation-verification findings related to Research Question 2. The RQ1 discussion integrates baseline pre-survey results and thematic analysis of interview transcripts to explain inventory tracking and financial reconciliation problems. The RQ2 discussion examines the developed architecture, access controls, transaction integrity, and transmission security through executed checks and live deployment inspection. Its technical outcomes are separate from respondent-based validation ratings. The chapter does not report post-interaction acceptability results for Research Question 3.')
anchor=next(p for p in d.paragraphs if p.text=='Summary of Findings')
section_start=len(d.paragraphs)
rq2_text=[]
def add(txt,h=False):
 rq2_text.append(('heading' if h else 'paragraph',txt));return para(anchor,txt,h)

add('Findings for Research Question 2',True)
add('Research Question 2 asks which system framework and information security measures are appropriate for the cloud-based inventory and financial management system in terms of architecture, data security, user privileges, and data integrity and transmission security. The executed assessment found a working centralized framework with effective controls in the tested workflows. However, two reproducible authorization failures prevent an unconditional security-validation pass. The findings below identify what was observed, the controls that worked, and the conditions that must be addressed before a stronger suitability conclusion can be drawn.')
add('Scope and Evidence of the Technical Assessment',True)
add('The AI-assisted assessment reviewed the current repository and the connected KodiGo database on 27 October 2026. TypeScript compilation, ESLint, and the production build completed successfully. All 17 existing unit tests passed. A separate set of 22 controlled database scenarios produced 20 passes and 2 failures. Four anonymous browser-route checks and four HTTPS or API checks also passed. These results describe specific executed checks; they do not establish a population estimate, a questionnaire rating, or complete security coverage.')
add('The database scenarios used synthetic Admin, Cashier, Inventory, Super Admin, and MFA-enrolled identities across two test stores. Every fixture write was enclosed in a subtransaction that was rolled back, and a post-check confirmed that no synthetic Auth identities remained. Role and JWT-claim simulation exercised PostgreSQL authorization and transaction behavior directly. Actual sign-in, real TOTP enrollment and verification, device integration, recovery procedures, sustained load, and concurrent-client stress were outside this assessment. The repository integration script was not executed because its dedicated service-role test credential was unavailable; the controlled SQL checks are reported separately.')
table(anchor,'Table 4.6. Executed Technical Verification Results',
 ['Evidence family','Result','Interpretation'],[
 ['Compilation and build','Passed','TypeScript, lint, and production build completed. Build-size warnings remain.'],
 ['Existing unit tests','17 of 17 passed','Reporting and consignment calculation tests; no user acceptability inference.'],
 ['Controlled database scenarios','20 of 22 passed','Two confirmed authorization failures require remediation.'],
 ['Anonymous browser routes','4 of 4 passed','Dashboard, inventory, POS, and Super Admin routes redirected to login.'],
 ['HTTPS and anonymous API checks','4 of 4 passed','TLS 1.3 certificate verification succeeded; protected endpoints returned HTTP 401.'],
 ['Client credential scan','105 files checked','No matching private key, secret-key, or service-role JWT format was detected in scanned client artifacts.'],
 ['Live database inspection','36 tables and 8 views','All 36 public tables had RLS enabled; all 8 public views used security_invoker=true.']
 ],[2.1,1.15,2.75])
add('Table 4.6 separates test families because they measure different properties. The 20 passing database scenarios and the 17 passing unit tests must not be treated as expert agreement or an acceptability score. Similarly, enabling Row Level Security indicates that a database control exists; it does not prove that every policy or privileged function is correct. Supabase documents RLS as the database layer for restricting row access and advises reviewing privileges and security-definer functions as part of the access model (Supabase, n.d.-b, n.d.-c).')

add('System Architecture and Framework Design',True)
add('The examined framework uses a React and TypeScript browser application with Supabase services and PostgreSQL as the authoritative data store. Source inspection confirmed route guards and client state management, while live inspection confirmed database policies, business functions, and deployed user-management and invite functions. Centralized transaction handling addresses the fragmented records identified in RQ1 by writing related sales, stock, payment, and receipt information through server-side workflows. Bounded browser queues support selected interrupted operations, but a queued item is not equivalent to a committed database transaction.')
add('The build and route checks show that the current application can compile and direct unauthenticated visitors away from protected pages. Large output chunks and mixed static and dynamic imports remained build warnings. These are maintenance and delivery concerns rather than observed authorization failures. The evidence supports the feasibility of the selected framework for the examined workflows, but no comparative evaluation of alternative architectures or scalability benchmark was conducted. Accordingly, the assessment does not establish that this is the most appropriate framework among all possible alternatives.')

add('Data Security and Access Control',True)
add('All 36 inspected public tables had Row Level Security enabled, and all eight public views used invoker security. In controlled tests, an Admin could read an assigned-store product but could not read the product in another test store. Anonymous product access was denied. The direct HTTPS product request and both privileged Edge Function requests also returned HTTP 401 without an authenticated session. Four local browser routes redirected anonymous visitors to the login page. These findings support the use of database authorization together with interface access checks.')
add('MFA assurance was effective in the tested store-membership helper: an enrolled synthetic account at AAL1 did not receive store authorization, whereas the same identity at AAL2 did. However, the create_store_with_owner function accepted a store-creation request for that enrolled account at AAL1. This confirms inconsistent MFA enforcement across the tested privileged operations. The failed scenario should be corrected by applying the intended assurance requirement at the server authorization boundary of store creation and reviewing comparable functions. Supabase distinguishes factor enrollment from assurance-based access enforcement; both are needed for the intended MFA behavior (Supabase, n.d.-a).')

add('Role Based Access and User Privilege Management',True)
add('The controlled role tests supported the intended separation of everyday work. Cashier could read an assigned product but could not edit it or promote its own role. Inventory could update an assigned product but could not process a POS sale. An unassigned Super Admin could neither read a store product nor process a POS sale. A cross-store sale request was rejected, and an Admin could not delete the tested profile in another store. These outcomes demonstrate authorization in the tested database paths, without implying that every screen and operation received equivalent coverage.')
add('A second failed scenario showed that the anonymous database role could call log_audit_event and insert a synthetic audit entry without an authenticated identity. This finding concerns the authenticity and controlled creation of audit evidence; it does not show that the caller could read or alter arbitrary business records. The function requires review of execution grants and caller validation. The Security Advisor also flagged 18 public security-definer functions executable by anonymous callers and 38 executable by signed-in callers. Those notices identify functions for review rather than 56 proven vulnerabilities, because an intentionally exposed function may still enforce appropriate checks internally.')
table(anchor,'Table 4.7. Confirmed Authorization Failures in the Controlled Assessment',
 ['Scenario','Priority','Observed failure and required action'],[
 ['DB17 MFA assurance for store creation','High','An enrolled account at AAL1 created a store. Enforce the intended MFA assurance in the privileged function and retest.'],
 ['DB20 Anonymous audit event creation','Medium','An anonymous caller inserted an audit event through an RPC. Restrict direct execution and validate callers before accepting audit writes.']
 ],[1.9,.7,3.4])
add('The priorities in Table 4.7 are reviewer judgments based on the observed impact, not CVSS scores or participant ratings. The first failure bypasses the intended second-factor requirement for a privileged operation. The second weakens the trustworthiness of the audit trail and permits unauthenticated event creation. No production security configuration or application source was changed during this assessment, so the findings describe the pre-remediation state tested on the stated date.')

add('Data Integrity and Transmission Security',True)
add('The checkout scenarios provided direct evidence of transaction validation. A request with altered totals was rejected and left no sale row or stock change. A valid two-piece sale reduced synthetic stock from ten to eight and produced one receipt. Repeating the same sale identifier returned a single recorded sale without a second stock deduction. Requests with insufficient stock or an unassigned store were rejected. A later selling-price change did not alter the historical sale-item price. These findings support server-side validation, retry handling, and historical snapshots for the scenarios exercised.')
add('The backend HTTPS connection negotiated TLS 1.3 with a certificate that passed verification at the time of the test. This demonstrates encrypted transmission for that inspected connection, not a complete audit of every supported cipher, hosting endpoint, or device. A targeted scan of 105 source, public, and generated client files detected no concrete private-key, Supabase secret-key, or service-role JWT format. That scan is limited to its defined patterns and inspected artifacts; it is not a guarantee that no credential exposure exists elsewhere.')

add('Remaining Security and Validation Conditions',True)
add('The Security Advisor reported ten functions with a mutable search path and disabled leaked-password protection. It also identified the invoice_counters table as having RLS enabled without a policy. The latter does not by itself establish exposure: absence of an applicable policy normally denies direct row access, and a counter may intentionally be maintained through a protected function. These notices require review against the intended design. Relevant Supabase guidance treats advisor results as checks to investigate rather than automatic proof that every flagged configuration is exploitable (Supabase, n.d.-c).')
add('Migration history and repository files were compared with live function definitions. The live MFA helper existed even though its migration version was not separately recorded in the inspected ledger. Conversely, the live refund function did not show the refunded_amount update present in the newer local refund migration. These observations show why deployment validation must inspect final definitions instead of assuming that either a baseline schema or a migration ledger fully represents runtime behavior. The refund difference was identified by definition inspection; its business effect was not tested in this exercise.')
add('The assessment did not conduct a restore drill, penetration test, comprehensive storage-policy review, or real-account MFA challenge. It also did not establish an expert-survey suitability rating. Further verification should reproduce the failed scenarios after remediation, extend negative tests to other privileged functions, validate deployed migration definitions, and retain dated evidence for the corrected state. These conditions delimit the conclusions that the current evidence can support.')

add('Technical Answer to Research Question 2',True)
add('For the examined Kodigo implementation, a centralized browser and cloud-database framework with authenticated sessions, role and store restrictions, database RLS, invoker-security views, server-side transaction checks, and assurance-based MFA provides a technically supported foundation for the diagnosed SME workflows. The executed checks demonstrate specific working controls for access separation, stock consistency, historical records, and encrypted API communication. However, the observed MFA bypass in store creation and unauthenticated audit-event insertion mean that the tested implementation does not receive an unconditional security-validation pass. Its technical suitability is conditional on remediation and retesting, while comparative architectural suitability and the expert-rating component of the original RQ2 require separate evidence.')

# Maintain the old RQ1 summary and add the accurately scoped RQ2 conclusion.
p=next(p for p in d.paragraphs if p.text.startswith('The quantitative and qualitative findings converge'))
replace(p,'For Research Question 1, '+p.text[0].lower()+p.text[1:])
ref=next(p for p in d.paragraphs if p.text=='References')
para(ref,'For the implementation-verification component related to Research Question 2, the framework compiled successfully and passed 17 existing unit tests, four anonymous route checks, and four HTTPS or API checks. Controlled database validation produced 20 passes and two failures across 22 scenarios. The evidence supports the tested role, store, transaction, and transmission controls, but inconsistent MFA enforcement in store creation and anonymous audit-event creation require correction. The overall technical conclusion is conditional suitability, with no fabricated expert ratings or user-acceptability results.')

# Append primary technical references without changing the existing literature.
refs=[
 'Supabase. (n.d.-a). Multi-factor authentication. Supabase Docs. Retrieved October 27, 2026, from https://supabase.com/docs/guides/auth/auth-mfa',
 'Supabase. (n.d.-b). Row level security. Supabase Docs. Retrieved October 27, 2026, from https://supabase.com/docs/guides/database/postgres/row-level-security',
 'Supabase. (n.d.-c). Advisors. Supabase Docs. Retrieved October 27, 2026, from https://supabase.com/docs/guides/observability/advisors'
]
# Place the entries in alphabetical order within the existing reference list.
ins=next((p for p in d.paragraphs if p.text.startswith('Unicommerce.')),d.paragraphs[-1])
for r in refs:
 p=para(ins,r);p.paragraph_format.line_spacing=1.5;p.paragraph_format.left_indent=Inches(.5);p.paragraph_format.first_line_indent=Inches(-.5)

# Request Word to refresh contents and page references on opening.
settings=d.settings.element
e=settings.find(qn('w:updateFields'))
if e is None:e=OxmlElement('w:updateFields');settings.append(e)
e.set(qn('w:val'),'true')
d.save(TARGET)

# Human-readable assessment with a reproducible, explicit test register.
report=['# Kodigo technical validation for RQ2','',
 'Assessment date: 27 October 2026 (Asia/Manila). Reviewer: Codex, an AI-assisted technical review. This report records executed verification and inspection, not human expert-survey ratings.','',
 '**Conclusion: conditional technical suitability; two confirmed authorization failures require remediation.**','',
 'The source thesis research design retains its separate qualified-evaluator survey. This review supplies implementation evidence and does not manufacture respondent participation, means, standard deviations, signatures, or acceptance scores.','',
 '## Executed checks','',
 '- TypeScript, lint and production build passed. Build size/import warnings remain.',
 '- Existing unit tests: 17 passed, zero failed.',
 '- Controlled live database scenarios: 20 passed, 2 failed. Synthetic identities and writes were rolled back; zero synthetic Auth identities remained.',
 '- Four anonymous browser routes redirected to login.',
 '- Backend TLS 1.3 and certificate check passed; three anonymous API/Edge requests returned 401.',
 '- Credential-pattern scan: 105 client files, no matching private/secret/service-role credentials.',
 '- Live metadata: RLS enabled on 36 public tables; invoker security on 8 public views.','',
 'Database tests simulate PostgreSQL roles and JWT claims. They do not validate signed JWT issuance, Auth login, actual TOTP challenge, or HTTP gateway behavior. The HTTPS tests are separate direct requests. The repository service-role integration harness was not run because its dedicated credential was unavailable.','',
 '## Controlled database scenario register','',
 '| ID | Expected behavior | Result | Observation |','|---|---|---|---|']
for t in tests:
 obs=t.get('observed') or t.get('error') or ', '.join(f'{k}={v}' for k,v in t.items() if k not in ['id','test','passed'])
 obs=re.sub(r'RQ2_VALIDATION_[a-f0-9]+','synthetic fixture',obs)
 report.append('| '+t['id']+' | '+t['test']+' | '+('PASS' if t['passed'] else 'FAIL')+' | '+obs.replace('|','/')+' |')
report+=['','## Confirmed failures','',
 'DB17 — High priority: an MFA-enrolled synthetic Admin with AAL1 could create a store through create_store_with_owner. Apply the intended MFA assurance requirement in that privileged function, review parallel functions, and rerun DB16–DB18.',
 '',
 'DB20 — Medium priority: the anonymous role could insert an audit event through log_audit_event. Review PUBLIC/anon execution privileges and enforce caller validation so that the audit trail cannot accept unauthenticated events. Retest both direct calls and legitimate nested business operations.',
 '',
 'Priorities are technical-review judgments, not CVSS scores. These were confirmed with rollback-only synthetic writes. No production remediation was performed.','',
 '## Security Advisor findings','',
 '| Advisor check | Findings | Interpretation |','|---|---:|---|']
for x in advisors:
 link=x['remediation'];report.append(f"| [{x['title']}]({link}) | {x['count']} | {x['level']}; review in the context of actual function bodies and access design. |")
report+=['','The 18 anonymous and 38 authenticated public security-definer execution notices are review targets, not a count of confirmed exploitable vulnerabilities. invoice_counters has RLS without policies; direct access is normally denied in the absence of an applicable policy.','',
 '## Deployment and scope limitations','',
 'The final live MFA helper exists despite its migration not being separately recorded in the retrieved ledger. The live refund function differs from the newer local refund migration concerning refunded_amount; this difference was inspected but not functionally exercised. Do not infer runtime state solely from migration filenames.',
 '',
 'Not assessed: full real-account sign-in/MFA enrollment and code verification; exhaustive storage and file-access tests; all authenticated UI workflows; load/concurrency stress; backup restoration; external penetration testing; expert-survey ratings; RQ3 acceptability.',
 '',
 '## Evidence and reproduction','',
 'Private evidence directory: `outputs/rq2-validation-2026-10-27/`. Key files: `run-checks.mjs`, `local-checks.json`, `controlled-security-checks.sql`, `controlled-security-results-v2.json`, `network-and-client-checks.mjs`, `network-and-client-results.json`, `live-metadata.json`, `function-metadata.json`, `live-initial-0.json`, `live-initial-1.json`, `live-initial-2.json`. No secret values are included in this report.',
 '',
 'Run local checks with the recorded Node runtime. The controlled SQL harness creates synthetic fixtures inside a subtransaction and deliberately rolls them back; reproduce it only against an authorized environment. Failed scenarios should be rerun after a reviewed fix and recorded as a new dated assessment, preserving this original baseline.',
 '',
 'Primary guidance: [MFA](https://supabase.com/docs/guides/auth/auth-mfa), [RLS](https://supabase.com/docs/guides/database/postgres/row-level-security), [Advisors](https://supabase.com/docs/guides/observability/advisors).','']
(O/'Kodigo_RQ2_Technical_Validation_Report.md').write_text('\n'.join(report),encoding='utf8')
(B/'rq2-section-text.json').write_text(json.dumps(rq2_text,ensure_ascii=False,indent=2),encoding='utf8')
(B/'authorship-and-change-record.json').write_text(json.dumps({'source_sha256':hashlib.sha256(SOURCE.read_bytes()).hexdigest(),'target':str(TARGET),'source_preserved':True,'assessment':'AI-assisted technical verification, no human ratings','controlled_cases':len(tests),'passed':20,'failed':2,'scope':['Chapter 3 implementation-verification supplement','Chapter 4 introduction','RQ2 findings and two summary tables','combined summary','three primary technical references']},indent=2),encoding='utf8')
print(TARGET)

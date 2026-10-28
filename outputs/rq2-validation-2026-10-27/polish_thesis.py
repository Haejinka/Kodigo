from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
P=Path(r'C:\Kodigo v0.1.0\deliverables\thesis\Chapter 1-4 Manuscript v3 RQ2 Technical Review.docx')
d=Document(P)
for p in d.paragraphs:
 if p.text.startswith('Table 4.6 separates') or p.text.startswith('The priorities in Table 4.7'):
  p.paragraph_format.space_before=Pt(8)
 if p.text.startswith('For the implementation-verification component related'):
  for _ in range(4):
   prev=p._p.getprevious()
   if prev is not None and prev.tag==qn('w:p') and not ''.join(prev.itertext()).strip():prev.getparent().remove(prev)
   else:break
  for r in p.runs:
   r.text=r.text.replace('The overall technical conclusion is conditional suitability, with no fabricated expert ratings or user-acceptability results.','The technical conclusion is conditional suitability. Expert-survey ratings and user acceptability remain separate evaluation strands.')
 if p.text=='References':
  p.paragraph_format.page_break_before=True;p.paragraph_format.keep_with_next=True
for p in d.paragraphs:
 if p.text.startswith('Supabase. (n.d.-'):
  p.paragraph_format.line_spacing=2;p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(0)
  p.paragraph_format.keep_together=True
d.save(P)
print('Adjusted table spacing, summary flow and reference pagination')

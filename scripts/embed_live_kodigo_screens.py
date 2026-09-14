from pathlib import Path
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

src = 'Kodigo_Governance_and_Full_Manual_v1.1_Revised.docx'
out = 'Kodigo_Governance_and_Full_Manual_v1.1_Revised.docx'
doc = Document(src)

def remove_paragraph(p):
    p._element.getparent().remove(p._element)

paras = doc.paragraphs
for p in paras:
    if p.text.startswith('Live interface validation was performed'):
        p.text = (
            'Live interface screenshots were captured from the current Kodigo build through the in-app browser on '
            '15 September 2026. Figures 3–8 are the approved live captures used to illustrate the documented '
            'Login, Dashboard, POS, Inventory, Suppliers, and Settings/Security workflows.'
        )
        p.style = doc.styles['Normal']

figure_map = [
    ('Figure 3. Kodigo Login Interface', 'deliverables/live_kodigo_screens/figure-03-login.png'),
    ('Figure 4. Kodigo Dashboard and Store Context', 'deliverables/live_kodigo_screens/figure-04-dashboard.png'),
    ('Figure 5. Kodigo POS Terminal and Product Search', 'deliverables/live_kodigo_screens/figure-05-pos.png'),
    ('Figure 6. Kodigo Product Management and Inventory', 'deliverables/live_kodigo_screens/figure-06-inventory.png'),
    ('Figure 7. Kodigo Supplier Management', 'deliverables/live_kodigo_screens/figure-07-suppliers.png'),
    ('Figure 8. Kodigo Settings and Account Security', 'deliverables/live_kodigo_screens/figure-08-settings-security.png'),
]

# Remove the old placeholder caption paragraphs.
for p in list(doc.paragraphs):
    if p.text.startswith('Figure ') and any(p.text == cap for cap, _ in figure_map):
        remove_paragraph(p)

# Insert captions and live screenshots immediately before Section 5.
before = next(p for p in doc.paragraphs if p.text.startswith('5. ACCEPTABLE USE AND BYOD POLICY'))
for caption, image_path in figure_map:
    cap = before.insert_paragraph_before(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    cap.runs[0].bold = True
    imgp = before.insert_paragraph_before()
    imgp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    imgp.paragraph_format.keep_with_next = False
    run = imgp.add_run()
    run.add_picture(image_path, width=Inches(5.9))

doc.save(out)
print(out)

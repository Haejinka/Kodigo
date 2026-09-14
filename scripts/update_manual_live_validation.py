from docx import Document
from docx.shared import Pt

src = 'deliverables/Kodigo_Governance_and_Full_Manual_v1.1_Revised.docx'
out = 'Kodigo_Governance_and_Full_Manual_v1.1_Revised.docx'
doc = Document(src)

for p in doc.paragraphs:
    if 'Version 1.1  |  14 September 2026' in p.text:
        for r in p.runs:
            r.text = r.text.replace('Version 1.1  |  14 September 2026', 'Version 1.1  |  15 September 2026')
    if p.text.startswith('Validated screenshots were not available in the supplied project files.'):
        p.text = (
            'Live interface validation was performed in the current Kodigo build through the in-app browser on '
            '15 September 2026. The captured screens confirmed the Dashboard at /dashboard, POS Terminal at /pos, '
            'Product Management at /inventory, and Settings at /settings, including the current store context, '
            'three displayed products, stock-status indicators, checkout search/cart layout, and store-branding fields. '
            'The labeled placeholders below identify the corresponding live captures to insert from the approved '
            'evidence set before publication; they are not representations of live screens.'
        )
        p.style = doc.styles['Normal']
    if p.text.startswith('Version 1.1: adviser feedback incorporated'):
        p.text = (
            'Version 1.1: adviser feedback incorporated for process diagrams, ITIL workflow, response-target rationale, '
            'TOC, live-build screenshot validation and screenshot placeholders, privacy notice, AUP obligations, '
            'BYOD acknowledgment, and figure/table numbering.'
        )

doc.save(out)
print(out)

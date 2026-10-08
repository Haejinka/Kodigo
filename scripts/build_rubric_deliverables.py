from __future__ import annotations

from datetime import date
from io import BytesIO
from pathlib import Path
import textwrap

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path.cwd()
OUT = ROOT / 'deliverables'
SCREEN = OUT / 'live_kodigo_screens'
VERSION = 'Version 1.3 | 5 October 2026'
CONFIRM = 'OWNER INPUT: '
DOCUMENT_OWNER = 'Evon Christian P Vergara'
NAVY = '17365D'
BLUE = 'DCEAF7'
PALE = 'F2F6FA'
MID = '5B6573'
GRAY = 'D9DEE5'
PAGE_WIDTH = 8.27
MARGIN = 0.65
TEXT_WIDTH = PAGE_WIDTH - 2 * MARGIN


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tc_pr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=90, start=105, bottom=90, end=105):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in('w:tcMar')
    if margins is None:
        margins = OxmlElement('w:tcMar')
        tc_pr.append(margins)
    for name, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = margins.find(qn(f'w:{name}'))
        if node is None:
            node = OxmlElement(f'w:{name}')
            margins.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tbl_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = f'w:{edge}'
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn('w:val'), 'single')
        node.set(qn('w:sz'), '5')
        node.set(qn('w:space'), '0')
        node.set(qn('w:color'), GRAY)


def mark_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    tr_pr.append(header)


def set_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement('w:cantSplit')
    tr_pr.append(node)


def set_run_font(run, name='Aptos', size=10, bold=None, color='20252B', italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), name)
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    display = OxmlElement('w:t')
    display.text = ' ' if instruction.startswith('PAGE') else 'Update this table in Word if it does not refresh automatically.'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instr, separate, display, end])
    set_run_font(run, size=8, color=MID)


def add_hyperlink(paragraph, text, url):
    relationship_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('r:id'), relationship_id)
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '0563C1')
    properties.append(color)
    underline = OxmlElement('w:u')
    underline.set(qn('w:val'), 'single')
    properties.append(underline)
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Aptos')
    fonts.set(qn('w:hAnsi'), 'Aptos')
    properties.append(fonts)
    run.append(properties)
    label = OxmlElement('w:t')
    label.text = text
    run.append(label)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def configure_doc(title, short_title, subtitle, with_toc=False):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(PAGE_WIDTH)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.72)
    sec.bottom_margin = Inches(0.68)
    sec.left_margin = Inches(MARGIN)
    sec.right_margin = Inches(MARGIN)
    sec.header_distance = Inches(0.32)
    sec.footer_distance = Inches(0.32)

    styles = doc.styles
    normal = styles['Normal']
    normal.font.name = 'Aptos'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(32, 37, 43)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in [('Title', 28), ('Heading 1', 17), ('Heading 2', 13), ('Heading 3', 11)]:
        style = styles[name]
        style.font.name = 'Aptos Display' if name in ('Title', 'Heading 1') else 'Aptos'
        style._element.rPr.rFonts.set(qn('w:ascii'), style.font.name)
        style._element.rPr.rFonts.set(qn('w:hAnsi'), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(5)
    title_ppr = styles['Title']._element.get_or_add_pPr()
    for border in list(title_ppr.findall(qn('w:pBdr'))):
        title_ppr.remove(border)
    for name in ('List Bullet', 'List Number'):
        styles[name].font.name = 'Aptos'
        styles[name]._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
        styles[name]._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
        styles[name].font.size = Pt(10)
        styles[name].paragraph_format.space_after = Pt(3)

    # Header/footer are deliberately restrained for print.
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = hp.add_run(f'KODIGO  |  {short_title.upper()}')
    set_run_font(r, size=8, bold=True, color=MID)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f'Kodigo Capstone Documentation  |  {VERSION}  |  Page ')
    set_run_font(r, size=8, color=MID)
    add_field(fp, 'PAGE')

    props = doc.core_properties
    props.title = title
    props.subject = subtitle
    props.author = DOCUMENT_OWNER
    props.last_modified_by = DOCUMENT_OWNER
    props.keywords = 'Kodigo, system design, governance, user manual, policy'

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(66)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run('KODIGO')
    set_run_font(r, name='Aptos Display', size=12, bold=True, color=NAVY)
    p = doc.add_paragraph(style='Title')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(title)
    set_run_font(r, name='Aptos Display', size=26, bold=True, color='000000')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    r = p.add_run(subtitle)
    set_run_font(r, size=12, color=MID)

    add_table(doc, 'Document Control', ['Control', 'Value'], [
        ('Version and date', VERSION),
        ('Status', 'Final owner-issued edition. Peer-review rubric audit complete: 19/25 (Good).'),
        ('Document owner', DOCUMENT_OWNER),
        ('Approval status', 'No separate approval required for this final document issue.'),
        ('Review cycle', 'Review at least annually and after material system, role, privacy, or workflow changes.'),
        ('Evidence basis', 'Repository implementation and supporting project documents reviewed on 5 October 2026. Live deployment facts are not claimed unless evidenced.'),
    ], widths=[1.65, TEXT_WIDTH - 1.65], table_no=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(18)
    r = p.add_run('Final owner-issued edition following the supplied peer-review rubric audit. Deployment-specific input fields are operational details, not document-approval gates.')
    set_run_font(r, size=9, italic=True, color=MID)

    doc.add_page_break()
    if with_toc:
        add_heading(doc, 'Contents', 1)
        p = doc.add_paragraph()
        add_field(p, 'TOC \\o "1-3" \\h \\z \\u')
        doc.add_page_break()
    settings = doc.settings.element
    update_fields = settings.find(qn('w:updateFields'))
    if update_fields is None:
        update_fields = OxmlElement('w:updateFields')
        settings.append(update_fields)
    update_fields.set(qn('w:val'), 'true')
    return doc


def add_heading(doc, text, level=1):
    return doc.add_paragraph(text, style=f'Heading {level}')


def add_para(doc, text='', lead=None, italic=False):
    p = doc.add_paragraph()
    if lead and text.startswith(lead):
        a = p.add_run(lead)
        set_run_font(a, bold=True)
        b = p.add_run(text[len(lead):])
        set_run_font(b, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_bullets(doc, items, numbered=False):
    for index, item in enumerate(items, start=1):
        if numbered:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.28)
            p.paragraph_format.first_line_indent = Inches(-0.28)
            p.paragraph_format.space_after = Pt(3)
            prefix = p.add_run(f'{index}. ')
            set_run_font(prefix, size=10)
        else:
            p = doc.add_paragraph(style='List Bullet')
        r = p.add_run(item)
        set_run_font(r, size=10)


def add_table(doc, caption, headers, rows, widths=None, table_no=None, font_size=8.7):
    if caption:
        label = f'Table {table_no}. {caption}' if table_no is not None else caption
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(label)
        set_run_font(r, size=9, bold=True, color='000000')
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if widths is None:
        widths = [TEXT_WIDTH / len(headers)] * len(headers)
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    header = table.rows[0]
    mark_repeat_header(header)
    set_no_split(header)
    for i, value in enumerate(headers):
        cell = header.cells[i]
        cell.width = Inches(widths[i])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        shade(cell, NAVY)
        set_cell_margins(cell)
        cell.text = ''
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(str(value))
        set_run_font(r, size=font_size, bold=True, color='FFFFFF')
    for row_i, values in enumerate(rows):
        row = table.add_row()
        set_no_split(row)
        for i, value in enumerate(values):
            cell = row.cells[i]
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            shade(cell, 'FFFFFF' if row_i % 2 == 0 else PALE)
            cell.text = ''
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    set_table_borders(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_figure(doc, caption, image, alt_text, figure_no, width=6.9):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f'Figure {figure_no}. {caption}')
    set_run_font(r, size=9, bold=True, color='000000')
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(9)
    pic = p.add_run().add_picture(image, width=Inches(width))
    pic._inline.docPr.set('descr', alt_text)
    pic._inline.docPr.set('title', caption)
    return p


def add_toc_heading(doc, section, title):
    return add_heading(doc, f'{section} {title}', 1)


def make_font(size, bold=False):
    paths = [
        'C:/Windows/Fonts/arialbd.ttf' if bold else 'C:/Windows/Fonts/arial.ttf',
        'C:/Windows/Fonts/calibrib.ttf' if bold else 'C:/Windows/Fonts/calibri.ttf',
    ]
    for path in paths:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def text_box(draw, box, title, lines, fill='EAF1F8', outline=NAVY, font_size=30):
    x1, y1, x2, y2 = box
    draw.rounded_rectangle(box, radius=24, fill='#' + fill, outline='#' + outline, width=4)
    title_font = make_font(font_size, True)
    body_font = make_font(max(18, font_size - 8))
    draw.text((x1 + 22, y1 + 18), title, font=title_font, fill='#172B4D')
    y = y1 + 62
    for line in lines:
        draw.text((x1 + 23, y), line, font=body_font, fill='#283443')
        y += max(25, font_size - 1)


def arrow(draw, start, end, color='#3976B9', width=7):
    draw.line([start, end], fill=color, width=width)
    import math
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    length = 21
    left = (end[0] - length * math.cos(angle - 0.55), end[1] - length * math.sin(angle - 0.55))
    right = (end[0] - length * math.cos(angle + 0.55), end[1] - length * math.sin(angle + 0.55))
    draw.polygon([end, left, right], fill=color)


def architecture_diagram():
    w, h = 2100, 1180
    im = Image.new('RGB', (w, h), 'white')
    d = ImageDraw.Draw(im)
    title = make_font(42, True)
    d.text((80, 35), 'Kodigo current logical architecture', font=title, fill='#17365D')
    # Actors
    text_box(d, (65, 180, 405, 490), 'Users', [
        'Admin', 'Cashier', 'Inventory', 'Super Admin',
    ], font_size=31)
    # Browser application and state
    text_box(d, (525, 150, 1055, 405), 'Browser application', [
        'React 19 + TypeScript + Vite', 'React Router pages and guards', 'Tailwind / Radix user interface',
    ], font_size=30)
    text_box(d, (525, 500, 1055, 750), 'Application state and services', [
        'Zustand stores and Supabase JS', 'IndexedDB cache and offline queues', 'Optional Web Serial cash drawer',
    ], fill='F4F7FA', font_size=29)
    # Supabase service boundary
    text_box(d, (1220, 145, 2015, 360), 'Supabase Auth', [
        'Email/password sessions', 'Optional TOTP factor and step-up checks',
    ], font_size=30)
    text_box(d, (1220, 415, 2015, 700), 'Postgres data services', [
        'Store membership and row-level security', 'Tables, views, RPCs and triggers', 'Audit, inventory and sale lifecycle records',
    ], fill='E8F2EC', font_size=29)
    text_box(d, (1220, 760, 2015, 1000), 'Supabase Edge Functions', [
        'admin-users managed account operations', 'generate-invite owner invitation codes', 'Server secrets stay outside browser code',
    ], fill='FFF4DE', font_size=27)
    arrow(d, (405, 325), (525, 275))
    arrow(d, (790, 405), (790, 500))
    arrow(d, (1055, 595), (1220, 285))
    arrow(d, (1055, 635), (1220, 550))
    arrow(d, (1055, 690), (1220, 860))
    # Cloud/static hosting status note
    d.rounded_rectangle((65, 830, 1055, 1000), radius=18, fill='#F5F5F5', outline='#A7AFB8', width=3)
    d.text((90, 850), 'Hosting and environment note', font=make_font(28, True), fill='#20252B')
    lines = [
        'Repository includes Vercel SPA rewrite configuration and local Vite dev server.',
        'Actual production host, domains, Supabase project, and applied migrations require team verification.',
    ]
    for i, line in enumerate(lines):
        d.text((92, 895 + i * 39), line, font=make_font(21), fill='#354052')
    stream = BytesIO()
    im.save(stream, format='PNG', dpi=(220, 220))
    stream.seek(0)
    return stream


def incident_diagram():
    w, h = 1550, 1850
    im = Image.new('RGB', (w, h), 'white')
    d = ImageDraw.Draw(im)
    d.text((70, 25), 'Incident management and escalation workflow', font=make_font(40, True), fill='#17365D')
    x1, x2 = 330, 1220
    box_h = 125
    nodes = [
        (95, 'Report and log', ['Record reporter, time, store, impact, evidence and current status.']),
        (270, 'Categorize and prioritize', ['Assess impact and urgency; assign P1 to P4 and an owner.']),
        (445, 'Tier 0 / Tier 1 diagnosis', ['User self-check, then Store Team Lead or Admin validates impact and safe workaround.']),
    ]
    for y, title_txt, lines in nodes:
        text_box(d, (x1, y, x2, y + box_h), title_txt, lines, font_size=29)
    arrow(d, (775, 220), (775, 270))
    arrow(d, (775, 395), (775, 445))
    # Resolved diamond
    diamond = [(775, 650), (925, 745), (775, 840), (625, 745)]
    d.polygon(diamond, fill='#E8F2EC', outline='#276749')
    d.line(diamond + [diamond[0]], fill='#276749', width=4)
    d.text((710, 728), 'Resolved?', font=make_font(27, True), fill='#17365D')
    arrow(d, (775, 570), (775, 650))
    d.text((948, 715), 'No', font=make_font(23, True), fill='#17365D')
    d.text((560, 715), 'Yes', font=make_font(23, True), fill='#17365D')
    # YES resolution path on left
    text_box(d, (55, 900, 575, 1085), 'Confirm and close', [
        'User validates service; document', 'resolution and close the incident.',
    ], fill='E8F2EC', font_size=25)
    arrow(d, (625, 745), (315, 900))
    # NO: Tier 2 then major decision
    text_box(d, (720, 900, 1370, 1085), 'Tier 2 technical investigation', [
        'System Administrator examines access, data,', 'sync, deployment and recovery evidence.',
    ], fill='EAF1F8', font_size=26)
    arrow(d, (925, 745), (1040, 900))
    diamond2 = [(1040, 1140), (1205, 1240), (1040, 1340), (875, 1240)]
    d.polygon(diamond2, fill='#FFF4DE', outline='#9A6700')
    d.line(diamond2 + [diamond2[0]], fill='#9A6700', width=4)
    d.text((943, 1215), 'Major, privacy,', font=make_font(22, True), fill='#17365D')
    d.text((960, 1245), 'or vendor issue?', font=make_font(22, True), fill='#17365D')
    arrow(d, (1040, 1085), (1040, 1140))
    d.text((1225, 1205), 'Yes', font=make_font(23, True), fill='#17365D')
    d.text((775, 1205), 'No', font=make_font(23, True), fill='#17365D')
    text_box(d, (1160, 1400, 1490, 1580), 'Tier 3 command', [
        'IT Manager leads major response;', 'Privacy Contact handles privacy', 'assessment; vendor as needed.',
    ], fill='FFF4DE', font_size=22)
    arrow(d, (1205, 1240), (1325, 1400))
    text_box(d, (630, 1400, 1110, 1580), 'Restore and validate', [
        'Recover service, reconcile records,', 'and confirm the result with reporter.',
    ], fill='E8F2EC', font_size=23)
    arrow(d, (875, 1240), (870, 1400))
    # Tier 3 rejoins restoration and closure; lessons loop to governance.
    arrow(d, (1160, 1490), (1110, 1490))
    text_box(d, (390, 1660, 1230, 1815), 'Record lessons and corrective action', [
        'Capture cause, evidence, owner and due date; feed approved changes into governance review.',
    ], fill='F4F7FA', font_size=25)
    arrow(d, (870, 1580), (810, 1660))
    # user confirmation path rejoins lessons
    arrow(d, (315, 1085), (315, 1735))
    arrow(d, (315, 1735), (390, 1735))
    stream = BytesIO()
    im.save(stream, format='PNG', dpi=(190, 190))
    stream.seek(0)
    return stream


def add_doc_footer_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run(text)
    set_run_font(r, size=8.7, italic=True, color=MID)


def build_aup():
    d = configure_doc('Kodigo Acceptable Use Policy', 'Acceptable Use Policy', 'Rules for authorized, secure, and accountable use of Kodigo')
    add_heading(d, '1. Purpose and policy statement', 1)
    add_para(d, 'This policy defines how authorized people may use Kodigo accounts, store data, reports, devices, networks, and connected equipment. It protects the confidentiality, accuracy, availability, and traceability of retail operations. Access is a work privilege limited to an approved role and assigned store.')
    add_heading(d, '2. Scope and definitions', 1)
    add_para(d, 'The policy applies to owners, administrators, cashiers, inventory staff, developers, technical support, reviewers, contractors, and other approved people who access Kodigo or its information in development, test, staging, production, exports, logs, or backups.')
    add_para(d, 'An assigned store is a branch linked to the user through an approved store membership. A restricted record is a credential, security record, payment reference, backup, or export that needs named, limited access. An incident is an event that may harm service availability, record accuracy, account security, or personal information.')
    add_heading(d, '3. Roles and duties', 1)
    add_table(d, 'AUP responsibility assignments', ['Role', 'Required responsibility'], [
        ('Every user', 'Use an individual account, follow store and role limits, protect records, complete required training, and report issues promptly.'),
        ('Admin / store owner', 'Approve operational access, maintain accurate store assignments, review local activity, and support orderly offboarding.'),
        ('System Administrator', 'Maintain technical services, protect privileged access, preserve logs, and investigate authorized incidents.'),
        ('IT Manager / policy owner', f'Approve policy exceptions and enforcement decisions. {CONFIRM}name the person or office that holds this authority.'),
        ('Privacy Contact', f'Coordinate privacy requests and privacy incident decisions. {CONFIRM}identify the organization Privacy Contact or DPO.'),
    ], [1.55, TEXT_WIDTH - 1.55], 2)
    add_heading(d, '4. Authorized use', 1)
    add_table(d, 'Permitted activity and obligation', ['Required behavior', 'Control obligation', 'Example'], [
        ('Perform assigned sales, inventory, procurement, reporting, administration, support, or review tasks.', 'Approved business use; least privilege.', 'A Cashier records a sale only in the assigned active store.'),
        ('Use only personal named credentials and approved authentication methods.', 'Accountability and credential protection.', 'Keep password, one-time code, reset link, and session private.'),
        ('Use approved browsers, networks, scanners, printers, and cash drawers.', 'Endpoint and availability control.', 'Use a compatible scanner that sends barcode input as keyboard text.'),
        ('Create a report or export only for an approved work purpose and keep it in an approved location.', 'Data minimization and controlled disclosure.', 'Save a stock report in an organization-managed folder with limited access.'),
        ('Use a test environment and test data for training and validation where available.', 'Environment separation.', 'Validate a release in staging before production approval.'),
        ('Correct an error through an approved adjustment, return, void, or new audit event.', 'Integrity and auditability.', 'Record the reason for a stock adjustment instead of silently changing a prior sale.'),
    ], [2.35, 1.65, TEXT_WIDTH - 4.0], 3)
    add_heading(d, '5. Prohibited activity', 1)
    add_table(d, 'Prohibited use and required control', ['Prohibited behavior', 'Obligation', 'Reason'], [
        ('Share accounts, passwords, reset links, OTPs, invite codes, API keys, or active sessions.', 'Identity and session control.', 'Preserves user attribution and prevents account takeover.'),
        ('View or change another store, role, user record, report, or log without authorization.', 'Store-scoped access control.', 'Prevents unauthorized disclosure or posting.'),
        ('Alter, delete, fabricate, conceal, or bypass evidence for sales, stock, users, receipts, closeouts, or audit events.', 'Integrity and auditability.', 'Keeps financial and inventory records reliable.'),
        ('Probe production, bypass a control, run an unapproved script, install an unapproved extension, or connect unknown equipment.', 'Change and endpoint security.', 'Reduces compromise and service interruption.'),
        ('Save or transmit full card numbers, PINs, OTPs, wallet passwords, or unnecessary personal information.', 'Payment-data safety and data minimization.', 'Reduces privacy and account risk.'),
        ('Upload Kodigo data to personal email, public storage, social media, public AI services, or removable media without written approval.', 'Approved disclosure and transfer.', 'Prevents uncontrolled copies and disclosure.'),
        ('Clear browser data while queued offline work remains, or deliberately disrupt synchronization, logging, or backups.', 'Continuity and evidence preservation.', 'Prevents loss, duplicate posting, or missing records.'),
    ], [2.55, 1.55, TEXT_WIDTH - 4.1], 4)
    add_heading(d, '6. Account, access, and resource rules', 1)
    add_bullets(d, [
        'Request access through the assigned Admin or approved access authority. State the business duty, role, store, start date, and any expiry date.',
        'The Admin verifies the requester and approves only the minimum role and store scope needed. Technical or privacy privileges require a separately recorded approval.',
        'Use the assigned account only. Lock or log out of a shared terminal when leaving it. Never let another person transact under your session.',
        'Use the All Stores view only for permitted aggregate review. POS checkout requires a specific active store.',
        'Do not use production data for demos, tests, screenshots, or training unless approved and minimized.',
        'The organization reviews roles and store memberships at least quarterly and immediately after termination, role change, transfer, or suspected compromise.',
    ], numbered=True)
    d.add_page_break()
    add_heading(d, '7. Data handling and privacy', 1)
    add_table(d, 'Minimum data handling rules', ['Data group', 'Required handling'], [
        ('Internal operational data', 'Use only for assigned work; do not publish externally.'),
        ('Confidential store, supplier, sales, stock, or personal data', 'Limit access by role and assigned store; use approved encrypted transfer and storage; restrict exports.'),
        ('Restricted credentials, security logs, payment references, or backups', 'Grant named need-to-know access, avoid casual export, preserve audit evidence, and dispose only under an approved retention schedule.'),
        ('Payment data', 'Keep only the supported payment method and limited or masked reference when needed. Kodigo must not store a full card number, PIN, OTP, or wallet password.'),
    ], [2.0, TEXT_WIDTH - 2.0], 5)
    add_para(d, 'The deploying organization is responsible for completing the privacy notice, deciding retention periods, responding to data-subject requests, and reviewing applicable Philippine privacy requirements. The repository does not establish the legal identity or contact details of the deploying organization.')
    add_heading(d, '8. Monitoring, training, and compliance', 1)
    add_para(d, 'Kodigo records application activity needed for operations, security, error diagnosis, transaction lifecycle, and audit evidence. Depending on configuration, records may include user and store context, transaction or stock events, report-export notices, client errors, authentication state, and limited browser or device context. Authorized administrators review these records only for a documented business, support, security, audit, or privacy purpose. This policy does not authorize indiscriminate inspection of personal content on a BYOD device.')
    add_bullets(d, [
        'Complete role-appropriate onboarding and refresher training before handling live transactions or sensitive administrative actions.',
        'The organization records the role, training date, relevant practice scenarios, trainer, and any remediation needed.',
        'The System Administrator reviews technical alerts and access evidence; the Admin reviews store-level access and operational discrepancies; the Privacy Contact reviews privacy-related events.',
        'The organization defines its support channel, record-retention period, and review owner before production deployment.',
    ])
    add_heading(d, '9. Incident reporting and response', 1)
    add_para(d, 'Immediately report suspected account compromise, lost device, unauthorized store access, incorrect or missing sale, unexpected stock movement, data exposure, suspicious activity, or a privacy concern through the approved support channel. Preserve the time, store, screen message, transaction or product identifier, and safe evidence. Do not include passwords, OTPs, full payment credentials, or unnecessary identity documents in a ticket.')
    add_para(d, f'Support channel: {CONFIRM}name the ticketing system, email, or phone channel and service hours.')
    add_para(d, 'The incident owner classifies impact and urgency, assigns a priority, documents actions, limits access where needed, escalates privacy or major incidents to the Privacy Contact and IT Manager, validates recovery with the reporter, and records lessons and follow-up actions. Only the authorized incident lead or Privacy Contact may coordinate external notifications.')
    d.add_page_break()
    add_heading(d, '10. Violations, enforcement, and exceptions', 1)
    add_para(d, 'A suspected violation is documented and assessed by an authorized manager. Where necessary to protect people or records, access may be temporarily restricted while the issue is investigated. The organization preserves relevant logs, informs the appropriate manager or Privacy Contact, gives the affected user a fair opportunity to provide facts under applicable organizational procedures, and records the outcome. Proportionate outcomes may include coaching, retraining, access correction, access suspension, disciplinary action, contract remedies, or lawful referral. Applicable employment, contract, and privacy requirements control.')
    add_para(d, 'An exception must state the business need, data and systems involved, compensating controls, accountable owner, approver, and expiry date. No exception may waive a legal requirement or authorize storing full card numbers, PINs, OTPs, or wallet passwords.')
    add_heading(d, '11. Document ownership and review', 1)
    add_para(d, f'Document owner: {DOCUMENT_OWNER}. This is the final issue following completion of the peer-review rubric audit; no separate document approval is required.')
    add_para(d, 'Review at least annually and when roles, system functions, relevant law, or risk conditions materially change. Keep approved versions and acknowledgments in a restricted, auditable location.')
    add_heading(d, '12. Acceptable Use Acknowledgment', 1)
    add_para(d, 'Complete before production access and again after a material policy revision. The organization retains the signed record under its approved retention schedule.')
    add_table(d, 'User acknowledgment form', ['Field', 'Entry'], [
        ('Full name', '____________________________________________________________'),
        ('User / employee ID', '____________________________________________________________'),
        ('Role and assigned store', '____________________________________________________________'),
        ('Policy version reviewed', VERSION),
        ('Training completed / date', '____________________________________________________________'),
        ('User signature / date', '____________________________________________________________'),
        ('Approver / supervisor / date', '____________________________________________________________'),
    ], [2.0, TEXT_WIDTH - 2.0], 6, font_size=9)
    add_para(d, 'Acknowledgment: I have read and understood this Acceptable Use Policy. I will use Kodigo only for authorized business purposes, protect my account, follow role and store access restrictions, handle data as required, report suspected incidents promptly, and complete required offboarding steps.')
    return d


def build_byod():
    d = configure_doc('Kodigo Bring Your Own Device Policy', 'BYOD Policy', 'Conditions for authorized access to Kodigo from personal devices')
    add_heading(d, '1. Purpose and scope', 1)
    add_para(d, 'This policy sets minimum conditions for accessing Kodigo from a personally owned phone, tablet, or computer. BYOD means bring your own device. A device is not approved merely because its browser can open the application. This policy supplements the Kodigo Acceptable Use Policy; both apply to BYOD access.')
    add_para(d, 'It applies to employees, owners, contractors, and support personnel whom the deploying organization has expressly authorized to use a personal device for Kodigo.')
    add_heading(d, '2. Eligibility and approval', 1)
    add_bullets(d, [
        'Use a personal device only after the Admin or designated approver records the user, device type, identifier when appropriate, role, store scope, business need, approval date, and review or expiry date.',
        'The organization may refuse or revoke BYOD access when a device is unsupported, compromised, shared in a way that exposes Kodigo data, or unable to meet this policy.',
        'Use a separate personal user profile for Kodigo work. Do not share a signed-in browser profile with family members or other unauthorized people.',
        f'Supported device models and browser versions: {CONFIRM}list tested operating systems, browser names, and minimum versions before granting BYOD access.',
    ], numbered=True)
    add_heading(d, '3. Device security requirements', 1)
    add_table(d, 'BYOD baseline controls', ['Requirement', 'Minimum rule', 'Purpose'], [
        ('Software support', 'Use a licensed, supported operating system and a current browser; enable automatic security updates.', 'Reduces exposure to known vulnerabilities.'),
        ('Screen lock', 'Use a strong passcode or password and automatic screen lock within five minutes. Biometrics may supplement, not replace, the passcode.', 'Prevents casual access to an unattended device.'),
        ('Encryption and protection', 'Enable full-device encryption, firewall, and active anti-malware where supported by the platform.', 'Protects locally stored information and network access.'),
        ('Device integrity', 'Do not use rooted or jailbroken devices, unsupported operating systems, unapproved extensions, or shared family profiles for Kodigo.', 'Preserves the integrity of the access environment.'),
        ('Network', 'Use a trusted network. Do not perform administrative or report-export work over public Wi-Fi unless an organization-approved protected connection is active.', 'Reduces interception and account risk.'),
        ('Physical handling', 'Keep the screen from public view, do not leave the device unattended while signed in, and lock it during breaks.', 'Protects information from nearby users.'),
    ], [1.35, 3.65, TEXT_WIDTH - 5.0], 2, font_size=8.3)
    add_heading(d, '4. Authentication and access control', 1)
    add_bullets(d, [
        'Use only your named Kodigo account. Do not share passwords, one-time codes, password-reset links, invite codes, or active sessions.',
        'Follow the role and assigned-store restrictions. A BYOD device does not grant broader access, and POS transactions still require a specific active store.',
        'Kodigo supports time-based one-time password (TOTP) multi-factor authentication and step-up checks in the repository migration set. The migration set does not establish which roles the live project requires to enroll.',
        'Do not save a password in an unmanaged browser or third-party password manager. Use only an organization-approved credential manager if one is provided.',
        'Sign out when finished on a shared or exposed device. Closing a tab alone may leave the session active.',
    ])
    add_para(d, f'MFA enrollment requirement by role: {CONFIRM}identify which roles require TOTP, any exception approver, and the target enforcement date.')
    add_heading(d, '5. Business data protection and offline use', 1)
    add_para(d, 'Kodigo uses browser storage for product lookup and eligible offline sales or selected mutations. This means business data can remain on a device while synchronization is pending. Treat browser storage and downloads as business records.')
    add_bullets(d, [
        'Do not download reports, receipts, screenshots, or user lists to a personal device unless a manager has approved the purpose, location, and retention period.',
        'Do not forward Kodigo records to personal email, public cloud drives, social networks, public AI services, or removable media without documented approval.',
        'When offline transactions are pending, keep the same browser profile and do not clear site data, use private browsing, uninstall the browser, or switch devices until synchronization succeeds or the work is handed over.',
        'After approved work is complete, remove unnecessary exports and screenshots from Downloads, email, chat, and other local folders using the device owner’s secure deletion method.',
        'Never enter or retain full card numbers, PINs, OTPs, or wallet passwords in a Kodigo record or personal note.',
    ])
    add_heading(d, '6. Lost, stolen, or compromised device', 1)
    add_para(d, 'Report a lost, stolen, shared, infected, or suspected compromised device immediately and no later than one hour after discovery. If the approved channel is unavailable, contact the supervisor or Admin by the alternate method listed below. Give the approximate time, last known store/session, device type, and whether offline work may be pending. Do not include the device unlock code in the report.')
    add_para(d, f'Primary incident channel: {CONFIRM}provide approved channel and service hours. Alternate contact: {CONFIRM}provide escalation contact.')
    add_bullets(d, [
        'The Admin or System Administrator verifies the identity of the reporter, removes or narrows account/store access as appropriate, and requests session revocation or password reset through supported controls.',
        'The incident owner checks for pending offline work before any browser data is cleared and reconciles receipts, sales, and stock after access is restored.',
        'Remote wipe or mobile-device-management capability is not evidenced in the repository. Do not promise remote erasure; use it only if the organization has separately deployed and approved that capability.',
        'The Privacy Contact assesses whether personal information may have been exposed and coordinates any required response.',
    ], numbered=True)
    add_heading(d, '7. Monitoring, privacy, and support boundaries', 1)
    add_para(d, 'The organization may review Kodigo account activity, authentication state, authorized store operations, audit events, report-export events, and limited error or browser context for support, security, compliance, or incident investigation. It does not gain authority under this policy to inspect unrelated personal files, messages, browsing history, photos, or personal accounts on the device.')
    add_para(d, 'The repository does not show full-device management, remote wipe, or continuous device telemetry. If the organization deploys such controls, it must document their scope, purpose, access, retention, notice, and privacy review before enrollment.')
    add_para(d, 'Support covers Kodigo configuration and supported browser/device integration. Personal hardware repairs, carrier service, home routers, unrelated applications, and personal backups remain the device owner’s responsibility. If an issue cannot be diagnosed safely on a personal device, support may ask the user to reproduce it on an approved managed device.')
    add_heading(d, '8. Device removal and offboarding', 1)
    add_bullets(d, [
        'The Admin removes or changes the user’s Kodigo role and store membership on the effective separation or role-change date.',
        'The user signs out, removes saved Kodigo credentials, and synchronizes or hands over any pending work before erasing browser data.',
        'The System Administrator revokes sessions or tokens where supported and records the completion of access removal.',
        'The user removes approved local exports, receipts, caches, and screenshots after confirming that required records are preserved in the authorized system.',
        'Before sale, recycling, or transfer of the device, the owner backs up lawful personal files and performs a secure factory reset. The organization does not direct deletion of personal content.',
    ], numbered=True)
    add_heading(d, '9. Exceptions, violations, and policy review', 1)
    add_para(d, f'Exceptions require a written business reason, device and data scope, compensating controls, accountable owner, approving authority, and expiry date. The Privacy Contact must review any exception involving personal information. {CONFIRM}identify the person authorized to approve BYOD and security exceptions.')
    add_para(d, 'A violation may result in immediate BYOD deauthorization, access restriction, retraining, or proportionate employment or contract action under applicable rules. The organization reviews this policy at least annually and after a material application or device-control change.')
    add_heading(d, '10. BYOD Acknowledgment Form', 1)
    add_para(d, 'Complete this form for each approved device. Approval is specific to the user, device, role, and store scope recorded below.')
    add_table(d, 'BYOD approval record', ['Field', 'Entry'], [
        ('Full name', '____________________________________________________________'),
        ('User / employee ID', '____________________________________________________________'),
        ('Role', '____________________________________________________________'),
        ('Assigned store / branch', '____________________________________________________________'),
        ('Device type', '____________________________________________________________'),
        ('Device identifier, if applicable', '____________________________________________________________'),
        ('BYOD approved', 'Yes / No / N/A'),
        ('Date policy reviewed', '____________________________________________________________'),
        ('Training completed', 'Yes / No  |  Date: __________________________________________'),
        ('Approver / supervisor', '____________________________________________________________'),
        ('Approval date / expiry date', '____________________________________________________________'),
        ('User signature / date', '____________________________________________________________'),
        ('Approver signature / date', '____________________________________________________________'),
    ], [2.35, TEXT_WIDTH - 2.35], 3, font_size=9)
    add_para(d, 'Acknowledgment: I understand that Kodigo access is limited to my approved role and assigned stores. I will protect this device and account, follow this BYOD Policy, report loss or suspected compromise promptly, preserve pending offline work, remove business data during offboarding, and cooperate with proportionate Kodigo access and security checks described in the policy.')
    return d


def build_manual():
    d = configure_doc('Kodigo User Manual', 'User Manual', 'Step-by-step operating instructions for Kodigo store teams', with_toc=True)
    add_heading(d, '1. About this manual', 1)
    add_para(d, 'This manual explains how store owners and staff use Kodigo for checkout, inventory, suppliers, purchasing, reporting, account access, and incident handling. Kodigo is a browser-based point-of-sale and store-operations application for sari-sari stores and small retail branches. It connects sales and inventory records so staff can reconcile transactions, monitor stock, and make replenishment decisions.')
    add_para(d, 'Use the role assigned to your account. If the screen, store, or available action differs from these steps, stop before posting a transaction and ask the Admin to verify your access. Screenshots in Figures 1–6 were captured from the project build on 15 September 2026; later interface releases may differ.')
    add_heading(d, '2. Roles and access', 1)
    add_table(d, 'Application roles and everyday duties', ['Application role', 'Main work', 'Access boundary'], [
        ('Admin', 'Store administration, staff accounts, product and supplier management, reports, analytics, notifications, settings, and sales actions where enabled.', 'Limited to assigned stores; approves operational changes and access.'),
        ('Cashier', 'Point-of-sale checkout, payment confirmation, receipts, and transaction handoff.', 'POS-focused; cannot use administrative modules.'),
        ('Inventory', 'Products, stock counts and adjustments, restocking and inventory-oriented reports.', 'Cannot access sales management, supplier administration, or general admin modules.'),
        ('Super Admin', 'Owner invite-code governance and onboarding.', 'Does not manage tenant store operations.'),
    ], [1.0, 3.0, TEXT_WIDTH - 4.0], 1, font_size=8.6)
    add_para(d, 'A technical support or Privacy Contact role is an organizational duty, not a fifth Kodigo application role. Ask your organization who holds those responsibilities.')
    add_heading(d, '3. Sign in and choose a store', 1)
    add_bullets(d, [
        'Open the organization’s approved Kodigo address and sign in with your assigned email and password.',
        'If you cannot sign in, use Forgot Password or request a managed password reset from the Admin. Support staff must never ask for your current password.',
        'After sign-in, check your name, role access, and the active store shown in the top bar. If you work in several branches, select the branch where you are working.',
        'The All Stores context is for authorized aggregate review only. POS checkout requires one specific active store.',
        'If no store appears or the wrong store is selected, do not create a sale, product, adjustment, or purchase order. Ask the Admin to correct the store membership.',
        'Use Logout when finished, especially on a shared terminal. Closing a browser tab alone may leave the session active.',
    ], numbered=True)
    add_para(d, 'Account Security is available to users for supported password and authenticator settings. Kodigo supports TOTP authenticator enrollment and verification where enabled by the organization. Follow the security instructions shown in the application; never share a one-time code.')
    add_heading(d, '4. Navigation and screen overview', 1)
    add_table(d, 'Main modules', ['Screen', 'What it is for', 'Typical result'], [
        ('Dashboard', 'Review daily sales indicators, best-selling products, recent transactions, and stock alerts.', 'A store-level operational summary.'),
        ('POS Terminal', 'Search or scan items, build a cart, take payment, and issue a receipt.', 'A completed sale or an explicitly pending offline sale.'),
        ('Product Management', 'Maintain product details, units, prices, active status, stock, movement history, restocking, and consignment views.', 'Updated catalog or documented inventory event.'),
        ('Suppliers', 'Maintain supplier records and view purchase and delivery details.', 'A supplier record linked to one or more authorized stores.'),
        ('Analytics and Rankings', 'Review time-based revenue, profit, category, hourly, transaction, and product-ranking data.', 'Filtered summaries based on the selected store and dates.'),
        ('Sales Reports', 'Generate sales and inventory reports and approved exports.', 'A workbook or PDF report, with export notification where configured.'),
        ('Settings', 'Manage store branding, assigned users, notification preferences, and account security.', 'Updated store or account configuration.'),
    ], [1.35, 3.2, TEXT_WIDTH - 4.55], 2, font_size=8.4)
    add_heading(d, '5. Process a sale', 1)
    add_para(d, 'Before checkout, confirm the active store and make sure the product, price, and quantity shown match the customer’s request. POS writes an online sale through the database sale-processing function. Eligible offline sales are stored in the browser queue for later replay.')
    add_bullets(d, [
        'Open POS Terminal. If prompted, select a specific active store; do not transact in an All Stores view.',
        'Find an item by name, SKU, or barcode. Press F2 to focus search. A supported USB or Bluetooth scanner that types barcode text may be used.',
        'Select the correct product or selling option. For unit or bulk options, verify the label, quantity, and displayed price. Stock is tracked against the product’s base units; bulk availability is derived from that stock.',
        'Review the cart line by line. Correct accidental items and quantities before charging.',
        'Choose Charge or press F9. Select the supported payment method. For cash, enter the cash received and verify the displayed change.',
        'For a digital payment method, enter only the supported method and limited reference requested by the screen. Never enter a full card number, PIN, OTP, or wallet password.',
        'Confirm the total, accept payment, and wait for the completion message. Do not repeat checkout because the screen is slow.',
        'Review the generated receipt. Print or reprint only under the organization’s procedure. The optional cash drawer requires compatible hardware, browser support, and permission.',
    ], numbered=True)
    add_heading(d, '5.1 Returns, refunds, voids, and closeout', 2)
    add_para(d, 'These actions affect cash and inventory. Open the transaction lifecycle or closeout action available to your role, find the original sale, verify the receipt and items, record a clear reason, and obtain any required Admin approval. Returned quantities and whether they return to stock must be verified before submission. Do not delete a sale to hide an error; use the supported void, refund, or return action so the transaction remains traceable.')
    add_heading(d, '6. Work during a connectivity interruption', 1)
    add_para(d, 'Kodigo may use a local product cache and IndexedDB queues for eligible sales and selected changes. An offline operation is not confirmed as synchronized until Kodigo reports successful replay.')
    add_bullets(d, [
        'Continue only if the screen explicitly allows offline work and the correct store is active.',
        'If the sale is shown as pending, do not repeat it. Note the time, terminal, and receipt or sale reference shown.',
        'Keep the browser profile intact. Do not clear browsing data, use private mode, switch profiles, or uninstall the browser while work is pending.',
        'Reconnect to the approved network and leave Kodigo open. Wait for the queue to replay and check for success or rejection messages.',
        'If a permission, database, or store-access error appears, stop retrying and contact support. Such errors may not be safe to queue as network failures.',
        'After synchronization, reconcile receipts and stock. Escalate any missing, duplicated, or rejected transaction.',
    ], numbered=True)
    add_heading(d, '7. Manage products and inventory', 1)
    add_heading(d, '7.1 Add or edit a product', 2)
    add_bullets(d, [
        'Open Product Management and select the specific store. Choose Add Product or open a product to edit it.',
        'Enter a unique SKU, product name, category, base selling unit, cost, selling price, starting stock, alert thresholds, and optional barcode or image.',
        'Where applicable, configure how the product is received from a supplier (for example, case or pack) and the conversion into base units. Check the conversion carefully before saving.',
        'Optional unit and bulk selling options define separate labels and prices but draw from the same base-unit stock. Review every conversion and bundle price before making it available at POS.',
        'Save and check that the product appears under the correct store with the expected stock and status.',
    ], numbered=True)
    add_heading(d, '7.2 Adjust or receive stock', 2)
    add_bullets(d, [
        'Count the physical stock before changing the system quantity.',
        'Open the product’s stock adjustment action. Record the amount, reason (such as restock, damaged, expired, lost, manual count, conversion, or other), and a useful note.',
        'For a supplier delivery, choose the receiving unit and confirm the conversion to base units, quantity, and total cost. Do not mark a purchase order received until goods and quantities have been checked.',
        'For consigned stock, identify the supplier who owns the goods, record required quantity and amount owed, and verify the ownership and supplier fields before saving.',
        'Review the before-and-after stock and movement/adjustment history. Ask the Admin to approve unusual or high-impact changes under local rules.',
    ], numbered=True)
    add_heading(d, '7.3 Archive products and review stock activity', 2)
    add_para(d, 'Archive a product when it should no longer appear in active inventory or POS. Products with sales history must be archived rather than hard-deleted so historical reports remain intact. Use the active/archived filter to locate records. The movement history and sales-velocity panels help explain stock changes and consumption; verify the store and date filters before acting on a recommendation.')
    add_heading(d, '8. Stock alerts, restocking, suppliers, and purchase orders', 1)
    add_bullets(d, [
        'Open Notifications or the Restocking tab to review low, critical, and out-of-stock alerts. Product or selling-option alerts may depend on configured thresholds.',
        'Confirm the physical count, pending offline transactions, and recent deliveries before preparing a reorder.',
        'Review the proposed quantity, base unit, supplier, lead time, recent consumption, and available purchasing unit. Suggestions are decision support, not automatic orders.',
        'Create a purchase order for the correct supplier and store. Check each product, quantity, unit cost, and expected delivery before sending.',
        'Use the supported Sent, Received, or Cancelled status. Select Received only after checking the delivered goods. Supplier score or history may refresh from the recorded order and delivery events.',
        'For a shared supplier, check which stores it serves and which products are linked to it. Do not expose or edit another store’s supplier data without authorization.',
        'The Consignment view displays supplier-owned stock and settlement balances where configured. Confirm payment and reconciliation outside Kodigo unless a deployed feature explicitly provides a settlement transaction.',
    ], numbered=True)
    add_heading(d, '9. Dashboard, analytics, rankings, and reports', 1)
    add_para(d, 'Dashboard shows daily revenue, transaction count, average order value, estimated profit, recent transactions, best-selling products, and stock alerts for the selected context. Analytics supports Today, 7-day, 30-day, and 90-day views with revenue/profit, hourly, category, and transaction information. Rankings aggregates sales by product for a selected period.')
    add_bullets(d, [
        'Confirm the active store or authorized aggregate before interpreting results.',
        'Choose the date range and filters, then refresh and review the totals before export.',
        'Check sale status filters so voided, refunded, or returned transactions are included or excluded as intended.',
        'Inventory users may access inventory-oriented reports; sales reporting and export access follows the assigned role.',
        'Store reports only in an approved location. Apply the organization’s retention period and delete obsolete copies securely.',
    ], numbered=True)
    add_heading(d, '10. Manage users, store settings, and account security', 1)
    add_heading(d, '10.1 User and store access', 2)
    add_bullets(d, [
        'Admins open Settings > User Management to list, invite/create, update, reset, or remove managed staff accounts through the server-side admin-users function where deployed.',
        'Before creating or changing an account, verify the person, role, store assignment, and approval. Assign the least access needed.',
        'Remove access promptly when a person leaves or changes roles. Review pending offline work before device/browser cleanup.',
        'Super Admin uses the invite-governance screen for owner registration codes and does not manage normal store operations.',
    ], numbered=True)
    add_heading(d, '10.2 Store branding and notifications', 2)
    add_para(d, 'In Settings > General, an Admin can maintain store name, registered/business name, tax registration and rate, TIN and branch details, terminal identifier, receipt text, contact details, address, and logo. Confirm legal and tax values with the deploying organization before saving. Notification preferences control available operational summaries; the actual options depend on the screen and deployment.')
    add_heading(d, '10.3 Password and authenticator settings', 2)
    add_bullets(d, [
        'Use Account Security or Settings > Security as available for your role to change your password through the authenticated account flow.',
        'To enable TOTP, open the two-factor authentication panel, scan the QR code in an authenticator app or enter the setup key privately, then enter the current six-digit code and choose Verify and enable.',
        'Treat the QR code, setup key, and one-time codes as secrets. Do not send them to support.',
        'Removing an authenticator factor may reduce protection and can change which protected actions are allowed. Follow the organization’s account recovery process if you lose access.',
    ], numbered=True)
    add_heading(d, '11. Troubleshooting', 1)
    add_table(d, 'Common issue and safe response', ['Problem', 'Likely cause', 'Safe response'], [
        ('Cannot sign in', 'Incorrect password, reset needed, missing profile or role.', 'Use Forgot Password; ask the Admin to verify the account and role. Do not share your password.'),
        ('No store or wrong store', 'Membership missing or stale active-store selection.', 'Stop before posting; refresh or sign in again; ask the Admin to verify assignment.'),
        ('POS unavailable', 'Role or device restriction, network issue, or service outage.', 'Check assigned role, active store, and network; capture the message and use the approved continuity process.'),
        ('Barcode not found', 'Wrong code, inactive/missing product, scanner mode issue.', 'Search by name or SKU; verify product and scanner input before creating anything.'),
        ('Sale appears pending', 'Offline queue or slow connection.', 'Do not repeat the sale. Reconnect, wait for sync, and reconcile using time and receipt reference.'),
        ('Stock is incorrect', 'Unreceived delivery, unsynchronized sale, or wrong adjustment.', 'Count physically; review sale, movement, adjustment, and PO history; record an approved correction.'),
        ('Report totals differ', 'Store/date/status filters or pending sync.', 'Confirm context and filters, wait for sync, compare transaction status, and send relevant IDs to support.'),
        ('Cash drawer fails', 'Unsupported browser, insecure context, missing permission, cable, or hardware issue.', 'Use compatible HTTPS/localhost, reconnect, grant permission only after an approved user action, and use manual cash handling.'),
        ('Access denied or database error', 'Role/store policy or database rule blocks the action.', 'Do not bypass or retry repeatedly. Verify assignment and escalate with the exact message.'),
        ('App appears stale', 'Cached data or inactive page.', 'Check connection, bring the tab forward, refresh when safe, and preserve pending offline data.'),
    ], [1.35, 2.05, TEXT_WIDTH - 3.4], 3, font_size=8.0)
    add_heading(d, '12. Ask for help and provide feedback', 1)
    add_para(d, 'Use the organization’s approved support channel. Include your name and role, store, date/time, screen, action, expected and actual result, error message, business impact, and relevant transaction/product reference. Redact unnecessary personal information from screenshots. Never include a password, OTP, full card number, or wallet password.')
    add_para(d, f'Support channel and service hours: {CONFIRM}publish current support contact and operating hours.')
    add_para(d, 'To suggest an improvement, describe the task, the point of difficulty, expected result, and how often it occurs. The organization should log and review user feedback through its approved support or feedback record; a dedicated feedback feature in Kodigo is not confirmed by the repository.')
    add_heading(d, '13. Interface screenshots', 1)
    add_para(d, 'These images are live captures from the project build dated 15 September 2026. They illustrate representative screens and sample data; labels, data, and layout may change in a later deployment. Figure 1 shows sign-in; Figure 2 shows the dashboard and active store; Figure 3 shows the POS search and cart; Figure 4 shows product management; Figure 5 shows supplier management; Figure 6 shows settings and account security.')
    screenshots = [
        ('Kodigo Login Interface', SCREEN / 'figure-03-login.png', 'Kodigo sign-in page with email and password fields.'),
        ('Dashboard and Active Store Context', SCREEN / 'figure-04-dashboard.png', 'Kodigo dashboard with selected store, sales indicators, recent transactions, best-selling products and stock alerts.'),
        ('POS Terminal and Product Search', SCREEN / 'figure-05-pos.png', 'Kodigo point-of-sale screen with product lookup, catalog cards, order summary and charge action.'),
        ('Product Management and Inventory', SCREEN / 'figure-06-inventory.png', 'Kodigo product management page with product list, stock status, prices, and management actions.'),
        ('Supplier Management', SCREEN / 'figure-07-suppliers.png', 'Kodigo supplier list with score, reliability, price, lead-time, and order columns.'),
        ('Settings and Account Security', SCREEN / 'figure-08-settings-security.png', 'Kodigo settings security screen showing password and authenticator controls.'),
    ]
    for i, (caption, path, alt) in enumerate(screenshots, 1):
        if i > 1:
            d.add_page_break()
        add_figure(d, caption, str(path), alt, i, width=6.9)
    add_heading(d, '14. Quick end-of-shift checklist', 1)
    add_bullets(d, [
        'Complete or cancel open carts; do not leave an ambiguous payment screen.',
        'Confirm pending offline transactions synchronized or hand them over to the next shift with references.',
        'Complete the approved cash closeout and record any variance.',
        'Secure receipts and exports; remove customer-visible information from the counter.',
        'Log out, lock the device, and hand off unresolved incidents with the support reference.',
    ])
    return d


def build_governance():
    d = configure_doc('Kodigo System Design and Governance', 'System Design and Governance', 'Technical design and accountable governance for the Kodigo capstone system', with_toc=True)
    add_heading(d, 'Executive summary', 1)
    add_para(d, 'Kodigo is a multi-store retail operations system for sari-sari stores and small branches. It combines browser-based checkout, catalog and stock controls, supplier and purchase-order workflows, reporting, and store-level access management. The design uses a React and TypeScript frontend with Supabase Auth and Postgres, database-enforced store scoping, server-side functions for privileged administration, and browser storage for eligible offline work. This document separates technical design from governance so reviewers can see both how the system works and who must control it.')
    add_para(d, 'This final document set records repository-verified design and governance requirements without presenting unverified live deployment settings as facts. The peer-review rubric audit is complete; deployment-specific owner inputs are identified in place and summarized in Section 24. They are operational details, not pending document approvals.')
    add_heading(d, 'PART I - SYSTEM DESIGN', 1)
    add_heading(d, '1. Purpose, business need, scope, and stakeholders', 1)
    add_para(d, 'Small retailers need checkout, stock records, supplier coordination, and simple performance reporting to stay aligned across daily work. When sales and inventory records are separated, cashiers may repeat a payment, staff may reorder from stale stock, and owners may be unable to trace a correction. Kodigo addresses this operational need by linking sale records to product stock, preserving transaction lifecycle events, showing stock alerts and restocking information, and summarizing activity by store.')
    add_table(d, 'System purpose and boundaries', ['Item', 'Description'], [
        ('Purpose', 'Support checkout, inventory accuracy, supplier procurement, and store-level review for small retail operations.'),
        ('Primary users', 'Admin, Cashier, Inventory, Super Admin; technical maintainers and reviewers have separate organizational duties.'),
        ('In scope', 'Browser application, Supabase Auth, Postgres data and policies, database functions and triggers, Edge Functions, IndexedDB offline queues, reports, and optional cash-drawer integration.'),
        ('Out of scope unless separately confirmed', 'Card processing, storing full payment credentials, a mobile device management service, automated legal/tax certification, a named production service desk, or verified production hosting and backups.'),
    ], [1.7, TEXT_WIDTH - 1.7], 2)
    add_heading(d, '2. Acceptance criteria and measurement', 1)
    add_para(d, 'The following are recommended acceptance tests and measurable targets, not reported test results. Record the test data, owner, environment, completion date, and evidence when the deployment team executes them; this document does not claim that unrun tests have passed.')
    add_table(d, 'Proposed success measures', ['Measure', 'Target for capstone acceptance', 'Evidence to retain'], [
        ('Role and store boundaries', 'All agreed route and database access tests pass for Admin, Cashier, Inventory, and Super Admin; every unauthorized cross-store read/write attempt is denied.', 'Dated test matrix and sanitized results.'),
        ('POS correctness', 'At least 20 representative checkout cases produce the expected totals, receipt references, and stock effects; zero duplicate sale records in the tested cases.', 'Staging test cases, transaction IDs, and stock reconciliation.'),
        ('Offline recovery', 'Five controlled offline sales replay once after reconnection; every test sale is either synchronized or clearly retained as an error for investigation.', 'Queue status evidence and reconciliation log.'),
        ('Inventory and purchasing', 'At least 10 product, adjustment, receiving, archive, and supplier scenarios match expected base-unit balances and history.', 'Scenario checklist and movement records.'),
        ('Recovery readiness', 'Complete one restore rehearsal in an isolated project and reconcile stores, products, sales, sale items, and adjustments before release approval.', 'Restore log, row-count reconciliation, and approver sign-off.'),
        ('User readiness', f'{CONFIRM}set number of representative users, training scenarios, and pass threshold.', 'Attendance, scenario results, trainer sign-off.'),
    ], [1.4, 3.7, TEXT_WIDTH - 5.1], 3, font_size=8.1)
    add_heading(d, '3. Architecture and major components', 1)
    add_para(d, 'The browser presents route-level screens and controls. Zustand stores manage session, active store, products, suppliers, cart, and alerts. The Supabase JavaScript client connects authenticated users to the backend. Supabase Auth issues sessions; Postgres policies constrain data; database RPCs and triggers carry business operations; Edge Functions perform privileged user and invite actions. IndexedDB holds a product cache and eligible offline queues. Web Serial is optional and limited to a compatible cash-drawer command.')
    add_figure(d, 'Kodigo Logical Architecture', architecture_diagram(), 'Architecture diagram showing users, browser application and state, IndexedDB, Supabase Auth, Postgres services, Edge Functions, and hosting verification note.', 1, width=6.95)
    add_heading(d, '4. Technology choices and rationale', 1)
    add_table(d, 'Technology stack and design reason', ['Layer', 'Current implementation', 'Reason for use / constraint'], [
        ('Frontend', 'React 19, TypeScript 5.9, Vite 7', 'Component-based interactive UI, typed application contracts, fast local development/build. TypeScript does not replace server authorization.'),
        ('Navigation and state', 'React Router 7, Zustand 5', 'Route organization and small focused stores for authentication, product, supplier, cart, and alert state.'),
        ('UI and charts', 'Tailwind CSS 4, Radix UI primitives, Lucide React, Recharts 3', 'Consistent responsive controls, accessible UI primitives, and interactive business charts.'),
        ('Backend', 'Supabase Auth, Postgres, row-level security, RPCs, triggers, views, Edge Functions', 'Managed sign-in and relational storage; database policies and functions enforce rules close to data. Actual target project and plan require confirmation.'),
        ('Offline', 'IndexedDB through idb', 'Browser-local product cache and queues support eligible work during connectivity interruptions; conflict resolution is limited and queues require careful reconciliation.'),
        ('Hardware and exports', 'Web Serial for optional ESC/POS drawer command; jsPDF and write-excel-file for exports', 'Uses browser capabilities for optional peripherals and report output. Printer/drawer compatibility must be tested.'),
        ('Hosting', 'Vercel SPA rewrite configuration exists in repository', f'Actual deployed host, domain, production project, CI release permissions, and hosting cost: {CONFIRM}confirm deployed configuration and owners.'),
    ], [1.05, 2.2, TEXT_WIDTH - 3.25], 4, font_size=8.1)
    add_heading(d, '5. Major modules and functional requirements', 1)
    add_table(d, 'Implemented module requirements', ['ID / module', 'Current behavior supported by repository', 'Expected result'], [
        ('FR-01 Identity and store context', 'Email/password sign-in, invite-based owner registration, role resolution, assigned store loading, and role-based route routing.', 'A user sees only authorized screens and store data.'),
        ('FR-02 POS and receipts', 'Search by name/SKU/barcode, cart, unit/bulk options, payment, receipt snapshot, and sales lifecycle actions through database functions.', 'Sale, sale items, payment details, and stock effects are recorded consistently.'),
        ('FR-03 Offline continuity', 'Cache products; queue eligible sales and selected mutations in IndexedDB; attempt replay on reconnect or page activity.', 'Pending/error state remains visible for reconciliation; business rule errors are not silently treated as network failures.'),
        ('FR-04 Product and inventory', 'Product/category maintenance, purchase and selling units, price history, stock adjustment and movement records, archive/restore, and consumption history.', 'Base stock and its history remain traceable; sold products are archived rather than deleted.'),
        ('FR-05 Suppliers and purchasing', 'Supplier records and store links, product supplier assignments, restocking options, purchase orders, receiving functions, supplier scoring, and consignment lots/balances.', 'Delivery and supplier-owned stock can be distinguished and reviewed.'),
        ('FR-06 Alerts and notifications', 'Stock and selling-option conditions create notifications with user read state; report-export and transaction events can be surfaced.', 'Authorized users can review operational notices for assigned context.'),
        ('FR-07 Reporting and analytics', 'Sales and inventory reports, Excel/PDF export paths, revenue/profit/category/hour charts, transaction views, rankings, and velocity.', 'Users review filtered output and export only within their role.'),
        ('FR-08 Administration', 'Store branding and settings, managed staff account operations through admin-users Edge Function, invite generation through generate-invite, password reset, TOTP management.', 'Authorized changes are attributed and limited to the required store scope.'),
    ], [1.35, 3.8, TEXT_WIDTH - 5.15], 5, font_size=7.9)
    add_heading(d, '6. Data model and relationships', 1)
    add_para(d, 'The schema is relational and store-oriented. Most operational records carry a store identifier or reach a store through an approved mapping. Row-level security (RLS) means database policies decide which rows an authenticated user can read or change for each request.')
    add_table(d, 'Core data entities', ['Domain', 'Entities evidenced in schema and migrations', 'Relationship / purpose'], [
        ('Identity and tenancy', '`profiles`, `stores`, `store_users`, `invite_codes`', 'Profiles map to stores through store_users; role and membership scope access.'),
        ('Catalog and stock', '`categories`, `products`, `product_selling_options`, `product_restocking_options`, `stock_adjustments`, `inventory_movements`, `inventory_stock_lots`', 'Products belong to store/category; stock movements explain sales, receipt, loss, adjustment, and consignment allocation.'),
        ('Sales', '`sales`, `sale_items`, `sale_payments`, `receipts`, `sale_events`, `sale_returns`, `sale_return_items`, `cashier_closeouts`', 'A sale has line items, payment and receipt records, with lifecycle events for void/refund/return and closeout.'),
        ('Procurement', '`suppliers`, `supplier_stores`, `product_suppliers`, `purchase_orders`, `purchase_order_items`', 'Supplier-store links support shared suppliers; orders link suppliers and products to receiving events.'),
        ('Operations and evidence', '`notifications`, `notification_user_states`, `audit_logs`, `error_logs`, `receipt_reprints`, `product_price_history`', 'Operational notices, user read status, errors, audit evidence, reprints, and price history support monitoring and traceability.'),
    ], [1.2, 3.1, TEXT_WIDTH - 4.3], 6, font_size=8.1)
    add_para(d, 'The list reflects repository migrations and application queries, not a verified live database inventory. The target project’s applied migration ledger and actual grants must be inspected before release.')
    add_heading(d, '7. Data flow and business process', 1)
    add_para(d, 'A named user signs in, receives a role and authorized store context, and performs an approved task. An online POS sale goes through the sale-processing RPC, which creates sale and item records and applies inventory effects. During an eligible network interruption, the browser may queue work locally and replay it later. Store staff reconcile any error or duplicate risk before closing the shift. Inventory movement, supplier delivery, alerts, reports, and review complete the operating cycle.')
    add_figure(d, 'Proposed Kodigo Business Process Flow', str(OUT / 'proposed_business_flow_diagram.png'), 'Eight-step business process flow from provisioning and authentication through catalog, sale, offline exceptions, replenishment, reporting, and governance.', 2, width=6.95)
    add_heading(d, '8. Role and permission design', 1)
    add_table(d, 'Application permission matrix', ['Capability', 'Admin', 'Cashier', 'Inventory', 'Super Admin'], [
        ('POS checkout', 'Execute where enabled', 'Execute', 'None', 'None'),
        ('Products and stock', 'Full store control', 'Read at POS', 'Manage within assigned scope', 'None'),
        ('Stock adjustment / receiving', 'Manage and approve per local controls', 'None', 'Record within role and store scope', 'None'),
        ('Suppliers and purchase orders', 'Manage assigned stores', 'None', 'No supplier administration', 'None'),
        ('Dashboard / analytics / rankings', 'Assigned-store view', 'None', 'Limited per route', 'None'),
        ('Reports', 'Sales and inventory', 'None', 'Inventory-oriented reports', 'None'),
        ('Store settings / staff management', 'Manage assigned stores via supported controls', 'Own account', 'Own account', 'None'),
        ('Invite governance', 'None', 'None', 'None', 'Generate owner invitation codes'),
        ('Account security', 'Own account', 'Own account', 'Own account', 'Own account'),
    ], [1.75, 1.3, 1.05, 1.45, TEXT_WIDTH - 5.55], 7, font_size=7.8)
    add_para(d, 'This summarizes implemented route intent and code paths. Database policies are the authority; route guards improve navigation but cannot replace backend enforcement. The exact live permission matrix must be validated against the deployed migration state.')
    add_heading(d, '9. Security architecture and risk controls', 1)
    add_table(d, 'Security controls classified by purpose', ['Control class', 'Control evidenced or required', 'Risk addressed / evidence'], [
        ('Preventive', 'Supabase Auth sessions; role checks; store membership; row-level security; authorized RPC checks.', 'Blocks unauthorized access and cross-store operations; validate with role/store test cases.'),
        ('Preventive', 'TOTP enrollment and assurance checks for protected operations where enabled; password reset flow.', 'Reduces account takeover; role-wide MFA requirement and target project settings require confirmation.'),
        ('Preventive', 'Server-side admin-users and generate-invite functions; service-role secrets stay in server environment.', 'Limits privileged user-management operations; verify only required secrets are configured and never expose them in VITE variables.'),
        ('Preventive', 'Store-specific active context for POS; database sale/receive/adjustment RPCs; validation and lifecycle rules.', 'Reduces wrong-store posting and direct stock mutation.'),
        ('Detective', '`audit_logs`, `sale_events`, `inventory_movements`, `error_logs`, closeouts, price history, and export notifications.', 'Supports accountability and incident reconstruction; retention, access review, and alert ownership must be assigned.'),
        ('Detective', 'Quarterly role/store review and release/test evidence.', 'Detects stale or excessive privileges and unsafe changes; retain dated approval records.'),
        ('Corrective', 'Revoke or change access, reset credentials, correct through traceable adjustment/return functions, and preserve error queues.', 'Contains compromise and repairs records without concealing prior events.'),
        ('Corrective', 'Production backup and restore procedure is described in the operations runbook.', 'Recovery capability depends on a configured backup plan and a successful restore drill; actual state must be verified.'),
    ], [1.05, 3.2, TEXT_WIDTH - 4.25], 8, font_size=7.7)
    add_heading(d, '10. Integrations, environments, and service identities', 1)
    add_table(d, 'Environment and identity inventory', ['Area', 'Repository evidence', 'Access / status', 'Owner input / evidence'], [
        ('Local development', 'Vite dev server in `kodigo-ui/`; `.env.example` names public Supabase URL and anon/publishable key.', 'Developer workstation; do not use production data for routine tests.', 'Document approved local data and credential handling.'),
        ('Test / staging', 'Integration harness expects an isolated project and test credentials; CI workflow is present.', 'Use only test users and non-production data; service-role key only in protected test environment.', 'Confirm project ID, owner, migration ledger, access list, and retention.'),
        ('Production', 'Application supports Supabase integration; Vercel SPA rewrite config is present.', 'Deployment existence, domain, plan, and applied migration version are not verifiable from repository files.', f'{CONFIRM}name production host/project/domain, owner, region, and applied migration list.'),
        ('Browser client', 'VITE_SUPABASE_URL and VITE_SUPABASE_ANON_KEY in example configuration.', 'Public client key may be delivered to browser; database policies must limit access.', 'Confirm only public keys are present in VITE build variables.'),
        ('Edge Function runtime', '`admin-users` and `generate-invite` functions; Supabase server environment includes service-role secret by design.', 'No human login sharing; secret remains in function environment and should be rotated under provider controls.', f'{CONFIRM}identify deployed function identities, secret owner, rotation schedule, and emergency access approver.'),
        ('Backup identity', 'Operations runbook recommends restricted backup storage and restore roles.', 'Actual backup subscription, vault, service account, and retention not verified.', f'{CONFIRM}name backup operator, restricted storage location, retention, and restore approver.'),
    ], [1.1, 2.55, 1.65, TEXT_WIDTH - 5.3], 9, font_size=7.25)
    add_para(d, 'Migration caution: the repository contains numbered migrations through migration_40 and additional dated migrations. The README setup sequence does not include every newest file. Treat the applied migration ledger in the target Supabase project as the deployment source of truth; do not assume repository files are already deployed or invent an execution order where the environment is unknown.')
    add_heading(d, 'PART II - SYSTEM GOVERNANCE', 1)
    add_heading(d, '11. Governance objectives and structure', 1)
    add_para(d, 'Governance assigns decision rights for access, changes, data, security, incidents, backups, and documentation. The deploying organization owns the service and accepts operational and privacy obligations. The roles below are governance assignments; only Admin, Cashier, Inventory, and Super Admin are application roles in the current TypeScript role union.')
    add_table(d, 'Governance responsibilities', ['Governance role', 'Owns / maintains', 'Approves / decides', 'Accountability'], [
        ('System Owner / deploying organization', 'Business purpose, user outcomes, service scope, policy and resource approval.', 'Accepts business changes, assigns owners, accepts residual risk.', 'Business impact, support capacity, and service direction.'),
        ('Admin / store owner', 'Store configuration, product and staff accuracy, local access, daily reconciliation.', 'Routine account/store access, stock corrections, and purchase actions under local thresholds.', 'Store operation and correct, timely escalation.'),
        ('Cashier', 'Accurate cart, payment, receipt, and shift handoff.', 'No role authority to approve own exceptions beyond assigned actions.', 'Actions performed under own credentials; prompt discrepancy report.'),
        ('Inventory staff', 'Catalog, physical counts, movement records, and inventory reports.', 'Routine stock activity within role; Admin approves unusual changes per policy.', 'Count accuracy, conversion correctness, and evidence notes.'),
        ('Super Admin', 'Owner invite-code governance.', 'Issue governed owner invitation codes.', 'Invite issuance and onboarding integrity; no tenant operations.'),
        ('System Administrator', 'Supabase configuration, technical access, releases, logs, and recovery work.', 'Execute approved changes; emergency action under incident authority.', 'Technical recovery evidence, privileged access, and migration records.'),
        ('IT Manager / service authority', f'Risk register, major incident coordination, release oversight. {CONFIRM}name accountable IT Manager or equivalent.', 'Approves major changes, emergency access, restoration, and service communications.', 'Service continuity and closure of corrective actions.'),
        ('Privacy Contact / DPO', f'Privacy requests, privacy assessments, notice and retention coordination. {CONFIRM}identify contact and organization.', 'Advises or decides privacy response within authority; coordinates any required notification.', 'Transparent, documented privacy handling.'),
        ('Developer / technical contributor', 'Source code and migrations under version control.', 'No routine production access; contributes reviewed changes.', 'Tests, review responses, and accurate technical documentation.'),
    ], [1.25, 2.25, 2.05, TEXT_WIDTH - 5.55], 10, font_size=7.35)
    add_heading(d, '12. Decision rights and segregation of duties', 1)
    add_table(d, 'Decision and approval matrix', ['Decision', 'Prepares', 'Approves', 'Evidence'], [
        ('Create or change a user/store assignment', 'Admin', 'Authorized store owner / manager', 'Named user, role, store, approver, date, expiry where relevant.'),
        ('High-impact stock correction', 'Inventory staff or Admin', 'Second authorized approver under local threshold', 'Count, reason, before/after, supporting reference.'),
        ('Normal code or database change', 'Developer', 'Reviewer and release owner', 'Issue, pull request, tests, migration, release record.'),
        ('Emergency technical change', 'System Administrator', 'IT Manager or delegated incident commander', 'Incident ID, reason, scope, rollback, time, retrospective review.'),
        ('Restore or backup-retention change', 'System Administrator', 'IT Manager / System Owner', 'Restore scope, approvals, reconciliation, retention evidence.'),
        ('Privacy request or disclosure', 'Privacy Contact with system support', 'Privacy authority under organization policy', 'Request, identity verification, decision, disclosure/correction/deletion evidence.'),
    ], [1.65, 1.35, 1.75, TEXT_WIDTH - 4.75], 11, font_size=7.8)
    add_para(d, 'Where project staffing is too small for full separation, the System Owner records the conflict, appoints an independent reviewer for high-risk actions, and retains evidence. No single person should silently request, approve, implement, and close their own high-impact production change.')
    add_heading(d, '13. Access governance lifecycle', 1)
    add_bullets(d, [
        'Request: document the user, business duty, role, store scope, start date, duration, and any privileged need.',
        'Approve: Admin or designated owner validates identity and need; an independent approver reviews high-risk or cross-store access.',
        'Provision: use the managed user function and invite flow selected for production; record the resulting role/store mapping.',
        'Review: at least quarterly, compare active users, roles, and store assignments with current duties; record decisions and exceptions.',
        'Revoke or change: remove access on separation, role change, compromise, or end of assignment; preserve transaction attribution and reconcile pending offline work.',
        'Privileged access: restrict to named administrators, use MFA where enforced, approve and time-bound support access, and review after use.',
    ], numbered=True)
    add_para(d, f'Production invite path: the repository UI currently calls the generate-invite Edge Function. Confirm the target project deployment, authorization policy, and administrator recovery path: {CONFIRM}name production invite owner and supported recovery procedure.')
    add_heading(d, '14. Data and privacy governance', 1)
    add_para(d, 'The deploying organization determines why personal information is collected, provides the required notice, assigns a Privacy Contact, approves retention, handles rights requests, and assesses incidents. Kodigo can contain names, email addresses, roles, store assignments, supplier contact details, sales and stock activity, payment method and limited reference, audit events, error context, exports, and device/browser details when linked to a user.')
    add_para(d, 'Kodigo must not store full card numbers, PINs, one-time passwords, or wallet passwords. Ordinary cash checkout does not require customer identity; collect customer details only where a defined and approved feature requires them.')
    add_para(d, 'Privacy handling should be implemented and reviewed in accordance with Republic Act No. 10173, the Data Privacy Act of 2012, its Implementing Rules and Regulations, and applicable National Privacy Commission issuances. This documentation is not a claim that Kodigo is legally certified or formally compliant.')
    reference = d.add_paragraph()
    reference.paragraph_format.space_after = Pt(4)
    lead = reference.add_run('Official reference texts (checked 5 October 2026): ')
    set_run_font(lead, size=9, bold=True)
    add_hyperlink(reference, 'Republic Act No. 10173', 'https://officialgazette.gov.ph/2012/08/15/republic-act-no-10173/')
    middle = reference.add_run('; ')
    set_run_font(middle, size=9)
    add_hyperlink(reference, 'National Privacy Commission Implementing Rules and Regulations', 'https://privacy.gov.ph/implementing-rules-regulations-data-privacy-act-2012/')
    add_table(d, 'Privacy governance fields to complete', ['Required field', 'Project-team action'], [
        ('Responsible organization', f'{CONFIRM}insert deploying organization legal name.'),
        ('Privacy Contact / DPO', f'{CONFIRM}insert name/title and authorized privacy contact.'),
        ('Official contact details', f'{CONFIRM}insert official email and business address.'),
        ('Retention schedule', f'{CONFIRM}approve retention periods by record class and legal/business need.'),
        ('Incident channel and notification authority', f'{CONFIRM}approve support channel, response owner, and external notification decision process.'),
    ], [2.0, TEXT_WIDTH - 2.0], 12)
    add_heading(d, '15. Risk management', 1)
    add_table(d, 'Initial project risk register', ['Risk and evidence', 'Potential effect', 'Control / owner', 'Open status'], [
        ('Deployed database may not match repository migration set; several new migrations and dated files exist.', 'Policies or columns expected by UI may be missing or different.', 'System Administrator compares applied ledger, policies, functions, grants and triggers before release.', 'Deployment state requires confirmation.'),
        ('Offline browser queues have limited conflict handling.', 'Duplicate, rejected, or stale changes may affect sales and stock.', 'Admin reconciles pending work; system owner defines offline eligibility; test queue replay and exceptions.', 'Resolve error items; do not clear browser data.'),
        ('Privileged operations depend on Edge Function secrets and deployment.', 'User administration or invite flow may fail or be misconfigured.', 'Restrict server secrets; verify function auth and role checks; use one documented production path.', 'Confirm deployment and secret ownership.'),
        ('MFA enforcement may differ by role or Supabase project settings.', 'Sensitive accounts could have inconsistent protection.', 'Confirm policy and required roles; test TOTP and step-up flows.', 'Deployment and policy confirmation.'),
        ('Backup plan is described in a runbook but actual subscription, schedule, and restore test are unverified.', 'Recovery may be unavailable or incomplete.', 'System Administrator verifies daily backup, 30-day restricted retention baseline, and restore drill; IT Manager approves.', 'Attach evidence before production acceptance.'),
        ('Policy owner, service channel, legal identity, and budget are not in repository evidence.', 'No clear accountability or resourcing for support and privacy response.', 'System Owner assigns accountable contacts, cost center, and response coverage.', 'Project team confirmation required.'),
        ('Generated screenshots and reports may expose store, user, or sales information.', 'Unnecessary disclosure in academic or public distribution.', 'Review/redact test data where needed, restrict distribution, and retain only approved evidence.', 'Review before publication.'),
    ], [2.0, 1.45, 2.55, TEXT_WIDTH - 6.0], 13, font_size=7.2)
    add_heading(d, '16. Change, release, and configuration management', 1)
    add_table(d, 'Repository and release standards', ['Area', 'Required standard'], [
        ('Source of truth', 'Latest applied SQL migration ledger first; then current runtime code and Edge Functions; then architecture and maintainer documents; baseline schema is not final RBAC state.'),
        ('Database changes', 'Add a new ordered migration; do not rewrite an already shared/applied migration. Review policy impact, data migration, rollback, and documentation.'),
        ('Naming and structure', 'Use descriptive migration filenames; snake_case database identifiers; keep feature UI in the existing components/pages/stores/lib folders; update role types, routes, navigation, RLS, validators and docs together.'),
        ('Environment values', 'Browser values use VITE_SUPABASE_URL and VITE_SUPABASE_ANON_KEY only. Never place service-role secrets in VITE variables, source control, or browser bundles.'),
        ('Review and testing', 'Use a reviewed change request; run npm ci, typecheck, lint and build; run integration tests only against an isolated test project with explicit opt-in.'),
        ('Deployment', 'Verify target environment and migrations, deploy required Edge Functions, inspect secrets, run smoke tests, and record release/rollback references. Vercel config is present; actual production setup requires confirmation.'),
        ('Rollback', 'Prefer frontend/function rollback for code issues; for schema/data damage, restore a verified backup to an isolated replacement and validate before promotion.'),
    ], [1.35, TEXT_WIDTH - 1.35], 14, font_size=8.0)
    add_heading(d, '17. Incident management and escalation', 1)
    add_para(d, 'The ITIL-aligned workflow logs, categorizes, assesses impact and urgency, assigns priority, diagnoses, escalates when needed, validates recovery, and records lessons. Tier 0 and Tier 1 handle self-check and store impact; Tier 2 investigates technical faults; Tier 3 coordinates major, privacy, vendor, or emergency decisions.')
    add_figure(d, 'ITIL-Aligned Incident Management and Escalation Process', incident_diagram(), 'Incident flowchart showing logging, prioritization, Tier 0 and Tier 1 diagnosis, resolved decision, Tier 2 investigation, major/privacy/vendor decision, Tier 3 command, restore, confirmation and lessons.', 3, width=6.65)
    add_table(d, 'Proposed response targets for capstone operations', ['Tier / priority', 'Trigger', 'Target / action', 'Owner'], [
        ('Tier 0', 'Single-user how-to or known warning; no security or data concern.', 'Immediate self-check; up to 10 minutes.', 'User.'),
        ('Tier 1', 'Persistent issue, cash discrepancy, one terminal/store, or access request.', 'Acknowledge P1/P2 within 15 minutes; P3 within one business hour; aim for workaround within one hour.', 'Store Team Lead / Admin.'),
        ('Tier 2', 'Multi-user issue, failed sync, suspected database/access defect, checkout outage, incorrect data.', 'Acknowledge within 30 minutes; P1 restore/workaround target 4 hours; P2 one business day.', 'System Administrator.'),
        ('Tier 3', 'Major outage, privacy/security event, restore/emergency change, vendor/legal impact.', 'P1 acknowledge within 15 minutes; update stakeholders every 60 minutes; incident commander sets recovery target.', 'IT Manager; Privacy Contact for privacy.'),
        ('P1 / P2 / P3 / P4', 'Priority reflects service and data impact; P4 is a request or planned change.', 'P1 immediate escalation; P2 validate/escalate quickly; P3 standard queue; P4 change/service-request process.', 'Assigned incident owner.'),
    ], [1.1, 2.5, 2.45, TEXT_WIDTH - 6.05], 15, font_size=7.45)
    add_para(d, 'These values are proposed capstone operating targets carried forward from the prior manual. The deploying organization sets its service hours, measurement rules, holidays, communication channel, and assigned owners in its operations runbook. These targets are not contractual service-level agreements.')
    add_heading(d, '18. Backup, recovery, audit, and monitoring responsibilities', 1)
    add_table(d, 'Operational ownership', ['Control activity', 'Responsible', 'Review / approval', 'Evidence'], [
        ('Daily production backup and restricted retention', 'System Administrator / provider operator.', 'IT Manager or System Owner.', 'Provider backup status, retention configuration, access list.'),
        ('Pre-migration backup and deployment', 'System Administrator.', 'Release approver.', 'Backup reference, migration log, test evidence.'),
        ('Restore rehearsal', 'System Administrator.', 'IT Manager verifies reconciliation and acceptance.', 'Restore test project, row-count checks, smoke-test record.'),
        ('Audit and error review', 'System Administrator and Admin within assigned scope.', 'IT Manager for elevated or major cases; Privacy Contact for personal data cases.', 'Dated review notes, incident links, corrective actions.'),
        ('Rollback / incident record', 'System Administrator or incident lead.', 'IT Manager / incident commander.', 'Timeline, impact, cause, evidence, recovery and due dates.'),
    ], [1.65, 1.75, 1.85, TEXT_WIDTH - 5.25], 16, font_size=7.7)
    add_para(d, 'The production operations runbook recommends daily Supabase backups, a manual backup before schema deployment, weekly exports of business-critical tables, restricted storage, at least 30 days of backup retention, and restore validation in a separate environment. These are the documented baseline recommendations; the deploying organization must verify that its subscription and configuration actually provide them.')
    add_heading(d, '19. User feedback and post-launch improvement', 1)
    add_para(d, 'The repository does not show a dedicated in-app feedback module. The System Owner uses the organization\'s support or feedback channel and maintains a feedback log with issue, role/store affected, frequency, business impact, evidence, decision, owner, and due date.')
    add_bullets(d, [
        'Collect user feedback during onboarding, support incidents, and a scheduled post-launch review.',
        'Group feedback into usability, reliability, data accuracy, access, and privacy themes; protect personal data in the record.',
        'Prioritize changes by operational harm, user frequency, legal/privacy impact, and implementation effort.',
        'Approve a change, assign an owner and due date, test in an isolated environment, and publish the release note or user instruction update.',
        'After release, ask the reporting users whether the change solved the problem and record any remaining issue.',
    ], numbered=True)
    add_para(d, f'Feedback channel and review cadence: {CONFIRM}name the support/feedback tool, accountable reviewer, and meeting schedule.')
    add_heading(d, '20. Documentation ownership, training, and review', 1)
    add_para(d, f'This document set is {VERSION}. Document owner: {DOCUMENT_OWNER}. The owner maintains alignment with the application, migration chain, and operating procedures. The peer-review rubric audit is complete; no separate document approval is required. Future technical or policy changes follow the normal change-control process.')
    add_bullets(d, [
        'Review at least annually and after material changes to roles, privacy processing, sale lifecycle, inventory controls, hosting, or incident ownership.',
        'Use version history in the project repository; record change summary, author/reviewer, approval, effective date, and replaced version.',
        'Keep controlled templates for access requests, risk reviews, incident records, restore tests, release approval, and user acknowledgment.',
        'Train Admins and Cashiers on sales and cash controls; Inventory staff on units, counts, and adjustments; privileged staff on access, MFA, privacy, and recovery.',
        'Retain attendance and practical competency evidence in an organization-approved restricted location.',
    ])
    add_heading(d, '21. Feasibility, rollout sequence, and resource plan', 1)
    add_para(d, 'The repository documents a working application, frontend build tools, Supabase integrations, and optional hardware integration. It does not include an approved capstone calendar, labor estimate, hosting budget, subscription tier, device count, or support staffing commitment. The staged plan below is a sequence of work gates; the project team must enter owners, dates, and cost values.')
    add_table(d, 'Indicative delivery sequence and owner inputs', ['Milestone', 'Completion evidence', 'Owner / target date'], [
        ('1. Confirm scope, owner, environment, and migration ledger', 'Approved scope, deployed-project inventory, role matrix and privacy contacts.', f'{CONFIRM}assign owner and date.'),
        ('2. Complete staging hardening and acceptance tests', 'Role/store tests, POS and offline evidence, restore drill, risk disposition.', f'{CONFIRM}assign owner, environment, and date.'),
        ('3. Train users and approve policies', 'Signed AUP/BYOD records, training attendance and competency checks.', f'{CONFIRM}set user count, trainer and date.'),
        ('4. Controlled release and operational handover', 'Release approval, backup evidence, support route, rollback and incident roster.', f'{CONFIRM}assign approver and date.'),
        ('5. Post-launch review', 'Feedback themes, metrics, incidents and prioritized improvement actions.', f'{CONFIRM}set review date and cadence.'),
    ], [2.0, 3.1, TEXT_WIDTH - 5.1], 17, font_size=7.7)
    add_table(d, 'Resource and budget inventory', ['Resource', 'Known project evidence', 'Cost / quantity to confirm'], [
        ('Software and development', 'React/Vite source, repository, npm toolchain, TypeScript, browser test environment.', f'{CONFIRM}confirm labor roles/hours and any paid developer tools.'),
        ('Cloud services', 'Supabase Auth/Postgres/Edge Functions and Vercel SPA rewrite configuration are used or referenced by code/config.', f'{CONFIRM}identify active provider plans, monthly cost, database/storage/egress assumptions.'),
        ('Devices and peripherals', 'Modern browser workstation; optional barcode scanner, printer, and Web Serial cash drawer.', f'{CONFIRM}list device count, existing vs purchased items, and unit/total cost.'),
        ('Training and support', 'Role-based user onboarding and incident handling are required by this governance plan.', f'{CONFIRM}assign trainer/support coverage and effort.'),
        ('Backup and recovery', 'Operations runbook recommends restricted backup storage and restore tests.', f'{CONFIRM}confirm plan cost, retention, storage size, and restore labor.'),
    ], [1.35, 3.0, TEXT_WIDTH - 4.35], 18, font_size=7.55)
    add_heading(d, '22. Course concept integration and evidence', 1)
    add_table(d, 'Course concepts applied to Kodigo', ['Concept area', 'Application in this design and governance', 'Evidence / section'], [
        ('Module 6: documentation ownership and version control', 'Named owner, version/date, controlled templates, scheduled review, and repository history.', 'Document Control; Sections 16 and 20.'),
        ('Modules 7-9: security principles and CIA', 'Least privilege, store scoping, named accounts, confidentiality of exports, integrity of sale/movement records, availability through queueing and recovery.', 'Sections 8, 9, 13 and 18.'),
        ('Modules 7-9: preventive, detective, corrective controls', 'Controls are explicitly classified and tied to account, record, and service risks.', 'Section 9.'),
        ('Module 10: user feedback and improvement loop', 'Collect, triage, approve, implement, release, and re-check user feedback after launch.', 'Section 19.'),
        ('Module 11: compliance and privacy responsibility', 'Data Privacy Act framework, responsible organization, privacy contact, rights requests, breach escalation, and no unsupported certification claim.', 'Section 14.'),
        ('Module 12: evidence-based proposal and readiness gates', 'Repository facts, measurable acceptance criteria, deployment-specific inputs, and operational readiness gates are explicit.', 'Sections 2, 10, 21 and 24.'),
    ], [1.65, 3.8, TEXT_WIDTH - 5.45], 19, font_size=7.75)
    add_heading(d, '23. Rubric traceability', 1)
    add_table(d, 'Peer review rubric criterion to document sections', ['Rubric criterion', 'Documents / sections carrying the requirement', 'Original assessment', 'Status after revision'], [
        ('1. Clarity and Completeness', 'System Design and Governance 1–2; User Manual 1–2; document controls in all four files.', 'Partial: purpose and users appeared; measurable outcomes and delivery ownership were incomplete.', 'Addressed. Remaining owner inputs are deployment details, not document-approval gates.'),
        ('2. System Design and Requirements', 'System Design and Governance 3–10; Figures 1–2; acceptance measures; user-facing flows in User Manual.', 'Partial: diagrams and modules existed, but technical architecture, current features and environment state were stale/incomplete.', 'Addressed from repository evidence; live environment and migration state are explicitly identified as owner-input evidence.'),
        ('3. Governance, Security and Compliance', 'System Design and Governance 11–20; Acceptable Use Policy 3–11; BYOD Policy 2–9.', 'Partial: RBAC, incidents and privacy notice existed; organizational authority, control classification, and enforcement detail were incomplete.', 'Addressed with explicit controls; named contacts and live backup/service configuration remain owner inputs.'),
        ('4. Course Concept Integration', 'System Design and Governance 17, 19–20, 22; AUP 4–10; BYOD 3–8.', 'Partial: COBIT/ITIL/CIA appeared but concepts did not consistently shape control design or post-launch feedback.', 'Explicitly mapped to controls and governance sections; course-specific module wording should be checked by adviser.'),
        ('5. Feasibility and Documentation Quality', 'All four controlled documents; System Design and Governance 21, 24; User Manual procedures and screenshots.', 'Partial: professional manual existed but combined policy/system content; dates, budget, schedule and technical caveats were absent or stale.', 'Four final deliverables; costs, staffing, calendar dates, and live production settings are labeled as owner inputs rather than assumed facts.'),
    ], [1.45, 2.15, 1.55, TEXT_WIDTH - 5.15], 20, font_size=7.2)
    d.add_page_break()
    add_heading(d, '24. Owner inputs for deployment-specific details', 1)
    add_bullets(d, [
        'Organization and accountability: legal organization name, System Owner, policy exception authority, IT service authority, Privacy Contact/DPO, and support channel.',
        'Database state: target Supabase migration ledger compared with repository migrations through migration_40 and dated additions; record applied state and drift.',
        'Environment and security: production/staging hosts and domains, Supabase projects, Edge Function deployment, public browser keys, and protected service-role secret owner.',
        'Operating controls: role-specific MFA settings, access-review owner, incident service hours and targets, backup retention, restore authority, and data-retention schedule.',
        'Rollout and resources: milestone dates, assigned people/hours, itemized provider/device/training costs, and funding source.',
        'Review live screenshots for any personal, supplier, or transaction data before wider academic or public distribution.',
    ])
    add_para(d, f'Final document status: peer-review rubric audit complete (19/25, Good). Document owner: {DOCUMENT_OWNER}. Section 24 lists deployment-specific owner inputs for operational records; it is not a document-approval request.')
    return d


def main():
    OUT.mkdir(exist_ok=True)
    paths = [
        (OUT / 'Kodigo_Acceptable_Use_Policy.docx', build_aup()),
        (OUT / 'Kodigo_BYOD_Policy.docx', build_byod()),
        (OUT / 'Kodigo_User_Manual.docx', build_manual()),
        (OUT / 'Kodigo_System_Design_and_Governance.docx', build_governance()),
    ]
    for path, doc in paths:
        doc.save(path)
        print(path)


if __name__ == '__main__':
    main()
